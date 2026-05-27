from __future__ import annotations

from datetime import datetime, timedelta
from pathlib import Path

from openpyxl import load_workbook

from src.campaign_service import CampaignService
from src.db import Database
from src.gui.campaign_view import build_live_send_confirmation_message


class FakeMailer:
    def __init__(self) -> None:
        self.sent: list[dict[str, object]] = []

    def send_email(self, **kwargs) -> None:
        self.sent.append(kwargs)


def make_db(tmp_path: Path) -> Database:
    db = Database(tmp_path / "outreach.sqlite")
    db.initialize()
    db.set_settings(
        {
            "sender_email": "sender@example.com",
            "daily_send_limit": "25",
            "delay_seconds": "0",
            "safe_mode": "true",
            "send_mode": "dry_run",
            "real_send_confirm_required": "true",
            "follow_up_delay_days": "2",
        }
    )
    return db


def add_approved_contact(db: Database, campaign_id: int, email: str) -> int:
    contact_id = db.add_contact(
        {
            "campaign_id": campaign_id,
            "email": email,
            "name": "Test",
            "company": "Outreach Automation",
            "topic": "Stage 1.4",
            "subject": "Stage 1.4 test",
            "base_message": "Controlled message",
            "status": "approved",
        }
    )
    db.update_contact(contact_id, {"generated_message": "Generated controlled message"})
    return contact_id


def test_dry_run_sent_can_be_re_approved(tmp_path: Path) -> None:
    db = make_db(tmp_path)
    campaign_id = db.get_default_campaign_id()
    contact_id = add_approved_contact(db, campaign_id, "owned@example.com")
    service = CampaignService(db, mailer=FakeMailer(), sleep_fn=lambda _: None)

    summary = service.send_approved(campaign_id)
    assert summary.dry_run_sent == 1
    assert db.get_contact(contact_id)["status"] == "dry_run_sent"

    assert service.approve_contacts([contact_id]) == 1
    assert db.get_contact(contact_id)["status"] == "approved"


def test_live_confirmation_message_contains_count_and_recipients() -> None:
    message = build_live_send_confirmation_message(
        6,
        [
            "one@example.com",
            "two@example.com",
            "three@example.com",
            "four@example.com",
            "five@example.com",
        ],
    )

    assert "Будет отправлено: 6" in message
    assert "one@example.com" in message
    assert "five@example.com" in message
    assert "...и еще 1" in message


def test_allowed_test_recipient_mismatch_error_includes_both_emails(tmp_path: Path) -> None:
    db = make_db(tmp_path)
    campaign_id = db.get_default_campaign_id()
    add_approved_contact(db, campaign_id, "blocked@example.com")
    db.set_settings(
        {
            "send_mode": "live",
            "allowed_test_recipient": "owned@example.com",
        }
    )
    service = CampaignService(db, mailer=FakeMailer(), sleep_fn=lambda _: None)

    summary = service.send_approved(campaign_id, confirm_live_send=True)

    assert summary.blocked_by_guardrail is True
    assert "owned@example.com" in summary.errors[0]
    assert "blocked@example.com" in summary.errors[0]


def test_daily_limit_error_includes_zero_remaining_count(tmp_path: Path) -> None:
    db = make_db(tmp_path)
    campaign_id = db.get_default_campaign_id()
    add_approved_contact(db, campaign_id, "owned@example.com")
    db.set_settings(
        {
            "send_mode": "live",
            "daily_send_limit": "1",
            "allowed_test_recipient": "owned@example.com",
        }
    )
    db.log_send(None, campaign_id, "send_email", "sent", "")
    service = CampaignService(db, mailer=FakeMailer(), sleep_fn=lambda _: None)

    summary = service.send_approved(campaign_id, confirm_live_send=True)

    assert summary.blocked_by_limit is True
    assert "Remaining today: 0/1" in summary.errors[0]


def test_follow_up_due_at_matches_configured_delay(tmp_path: Path) -> None:
    db = make_db(tmp_path)
    campaign_id = db.get_default_campaign_id()
    contact_id = add_approved_contact(db, campaign_id, "owned@example.com")
    db.set_settings(
        {
            "send_mode": "live",
            "daily_send_limit": "1",
            "allowed_test_recipient": "owned@example.com",
            "follow_up_delay_days": "2",
        }
    )
    service = CampaignService(db, mailer=FakeMailer(), sleep_fn=lambda _: None)

    summary = service.send_approved(campaign_id, confirm_live_send=True)
    contact = db.get_contact(contact_id)
    sent_at = datetime.fromisoformat(contact["sent_at"])
    follow_up_due_at = datetime.fromisoformat(contact["follow_up_due_at"])

    assert summary.sent == 1
    assert follow_up_due_at - sent_at == timedelta(days=2)


def test_export_report_includes_dry_run_sent_and_sent_statuses(tmp_path: Path) -> None:
    db = make_db(tmp_path)
    campaign_id = db.get_default_campaign_id()
    dry_run_contact_id = add_approved_contact(db, campaign_id, "dryrun@example.com")
    service = CampaignService(
        db,
        mailer=FakeMailer(),
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )

    dry_run_summary = service.send_approved(campaign_id)
    assert dry_run_summary.dry_run_sent == 1
    assert db.get_contact(dry_run_contact_id)["status"] == "dry_run_sent"

    add_approved_contact(db, campaign_id, "owned@example.com")
    db.set_settings(
        {
            "send_mode": "live",
            "daily_send_limit": "25",
            "allowed_test_recipient": "owned@example.com",
        }
    )
    live_summary = service.send_approved(campaign_id, confirm_live_send=True)
    assert live_summary.sent == 1

    report_path = service.export_report(campaign_id)
    workbook = load_workbook(report_path)
    rows = list(workbook["contacts"].iter_rows(values_only=True))
    headers = list(rows[0])
    status_index = headers.index("status")
    statuses = {row[status_index] for row in rows[1:]}

    assert "dry_run_sent" in statuses
    assert "sent" in statuses
