from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from src.background_worker import QueueJobProcessor
from src.campaign_service import CampaignService
from src.db import Database


class FakeMailer:
    def __init__(self, fail: bool = False) -> None:
        self.fail = fail
        self.sent: list[dict[str, object]] = []

    def send_email(self, **kwargs) -> None:
        self.sent.append(kwargs)
        if self.fail:
            raise RuntimeError("controlled fake smtp failure")


def write_smoke_csv(path: Path) -> None:
    path.write_text(
        "\n".join(
            [
                "email,name,company,topic,subject,message",
                "valid1@example.com,Valid One,Example Co,Queue Smoke,Hello {{company}},Message for {{name}}",
                "valid2@example.com,Valid Two,Example Co,Queue Smoke,Hello {{company}},Message for {{name}}",
                "valid3@example.com,Valid Three,Example Co,Queue Smoke,Hello {{company}},Message for {{name}}",
            ]
        ),
        encoding="utf-8",
    )


def make_service(tmp_path: Path, mailer: FakeMailer | None = None) -> CampaignService:
    db = Database(tmp_path / "outreach.sqlite")
    db.initialize()
    db.set_settings(
        {
            "send_mode": "dry_run",
            "safe_mode": "true",
            "real_send_confirm_required": "true",
            "daily_send_limit": "25",
            "delay_seconds": "0",
            "sender_email": "sender@example.com",
            "allowed_test_recipient": "",
        }
    )
    return CampaignService(
        db,
        mailer=mailer or FakeMailer(),
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )


def add_approved_contact(
    service: CampaignService,
    campaign_id: int,
    email: str,
) -> int:
    contact_id = service.db.add_contact(
        {
            "campaign_id": campaign_id,
            "email": email,
            "name": "Queue",
            "company": "Example",
            "topic": "Smoke",
            "subject": "Queue smoke",
            "base_message": "Queue smoke body",
            "status": "new",
        }
    )
    service.db.update_contact(
        contact_id,
        {"status": "approved", "generated_message": "Generated queue smoke body"},
    )
    return contact_id


def test_stage_1_6_1_full_queue_smoke(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    campaign_id = service.db.create_campaign("Stage 1.6.1 Queue Smoke")
    csv_path = tmp_path / "stage_1_6_1_queue_smoke.csv"
    write_smoke_csv(csv_path)

    import_result = service.import_file(csv_path, campaign_id)
    assert import_result.imported_count == 3
    assert import_result.skipped_count == 0
    assert {contact["status"] for contact in service.contacts(campaign_id)} == {"new"}

    generate_result = service.enqueue_generate_messages(campaign_id)
    assert generate_result.count == 3
    assert {contact["status"] for contact in service.contacts(campaign_id)} == {"generation_queued"}

    generated_jobs = QueueJobProcessor(service).process_available()
    assert generated_jobs == 3
    contacts = service.contacts(campaign_id)
    assert {contact["status"] for contact in contacts} == {"pending_review"}
    assert all(contact["generated_message"] for contact in contacts)
    assert {job["status"] for job in service.queue.list_recent_jobs(campaign_id)} == {"completed"}

    approved_target = sorted(contacts, key=lambda contact: contact["email"])[0]
    assert service.approve_contacts([approved_target["id"]]) == 1
    assert service.db.get_contact(approved_target["id"])["status"] == "approved"

    send_mailer = FakeMailer()
    send_service = CampaignService(
        service.db,
        mailer=send_mailer,
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )
    dry_run_result = send_service.enqueue_send_approved(campaign_id, mode="dry_run")
    assert dry_run_result.count == 1
    assert service.db.get_contact(approved_target["id"])["status"] == "queued"

    QueueJobProcessor(send_service).process_available()
    assert send_mailer.sent == []
    assert service.db.get_contact(approved_target["id"])["status"] == "dry_run_sent"
    assert service.db.fetch_one(
        "SELECT id FROM send_logs WHERE action = 'dry_run_send' AND contact_id = ?",
        (approved_target["id"],),
    )

    fail_campaign_id = service.db.create_campaign("Stage 1.6.1 Failing Job")
    fail_contact_id = add_approved_contact(service, fail_campaign_id, "fail@example.com")
    service.db.set_settings(
        {
            "send_mode": "live",
            "allowed_test_recipient": "fail@example.com",
            "daily_send_limit": "25",
            "real_send_confirm_required": "true",
        }
    )
    failing_service = CampaignService(
        service.db,
        mailer=FakeMailer(fail=True),
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )
    failing_service.enqueue_send_approved(
        fail_campaign_id,
        mode="live",
        confirm_live_send=True,
    )
    QueueJobProcessor(failing_service).process_available()
    failed_job = failing_service.queue.list_recent_jobs(fail_campaign_id)[0]
    assert failed_job["status"] == "failed"
    assert service.db.get_contact(fail_contact_id)["status"] == "failed"

    retry = failing_service.queue.retry_failed_job(failed_job["id"])
    assert retry.ok
    assert service.db.get_contact(fail_contact_id)["status"] == "queued"
    ok_mailer = FakeMailer()
    ok_service = CampaignService(
        service.db,
        mailer=ok_mailer,
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )
    QueueJobProcessor(ok_service).process_available()
    assert len(ok_mailer.sent) == 1
    assert service.db.get_contact(fail_contact_id)["status"] == "sent"

    cancel_campaign_id = service.db.create_campaign("Stage 1.6.1 Cancel Job")
    cancel_contact_id = add_approved_contact(service, cancel_campaign_id, "cancel@example.com")
    service.db.set_settings({"send_mode": "dry_run", "allowed_test_recipient": ""})
    cancel_service = CampaignService(
        service.db,
        mailer=FakeMailer(),
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )
    cancel_service.enqueue_send_approved(cancel_campaign_id, mode="dry_run")
    cancel_job = cancel_service.queue.list_recent_jobs(cancel_campaign_id)[0]
    assert cancel_service.queue.cancel_job(cancel_job["id"]).ok
    QueueJobProcessor(cancel_service).process_available()
    assert service.db.get_contact(cancel_contact_id)["status"] == "cancelled"
    assert service.db.fetch_one(
        "SELECT status FROM job_queue WHERE id = ?",
        (cancel_job["id"],),
    )["status"] == "cancelled"

    export_result = service.enqueue_export_report(campaign_id)
    assert export_result.count == 1
    QueueJobProcessor(service).process_available()
    reports = sorted((tmp_path / "exports").glob("contacts_export_*.xlsx"))
    assert reports
    workbook = load_workbook(reports[-1], read_only=True)
    assert {"contacts", "send_logs", "job_queue", "errors"}.issubset(workbook.sheetnames)
