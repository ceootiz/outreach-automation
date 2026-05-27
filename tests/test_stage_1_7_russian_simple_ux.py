from __future__ import annotations

import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication, QMessageBox

from src.background_worker import QueueJobProcessor
from src.campaign_service import CampaignService
from src.db import Database
from src.gui.i18n import status_to_ru, user_safe_error
from src.gui.main_window import MainWindow
from src.gui.settings_view import SettingsView
from src.mailer import ConnectionCheckResult


class FakeMailer:
    def __init__(self) -> None:
        self.sent: list[dict[str, object]] = []

    def send_email(self, **kwargs) -> None:
        self.sent.append(kwargs)


@pytest.fixture(scope="session")
def qapp() -> QApplication:
    app = QApplication.instance()
    if app is None:
        app = QApplication([])
    return app


@pytest.fixture(autouse=True)
def cleanup_qt_widgets(qapp: QApplication):
    yield
    for widget in QApplication.topLevelWidgets():
        campaign_view = getattr(widget, "campaign_view", None)
        if campaign_view is not None:
            campaign_view.shutdown()
        settings_view = getattr(widget, "settings_view", None)
        if settings_view is not None:
            settings_view.shutdown()
        if isinstance(widget, SettingsView):
            widget.shutdown()
        widget.close()
        widget.deleteLater()
    qapp.processEvents()


@pytest.fixture()
def quiet_message_boxes(monkeypatch: pytest.MonkeyPatch) -> list[tuple[str, str]]:
    messages: list[tuple[str, str]] = []

    def record(_parent, title: str, text: str, *args, **kwargs):
        messages.append((title, text))
        return QMessageBox.StandardButton.Ok

    monkeypatch.setattr(QMessageBox, "information", record)
    monkeypatch.setattr(QMessageBox, "warning", record)
    monkeypatch.setattr(QMessageBox, "critical", record)
    return messages


def make_service(tmp_path: Path) -> CampaignService:
    db = Database(tmp_path / "outreach.sqlite")
    db.initialize()
    db.set_settings(
        {
            "send_mode": "dry_run",
            "safe_mode": "true",
            "real_send_confirm_required": "true",
            "daily_send_limit": "25",
            "delay_seconds": "0",
            "sender_email": "",
            "allowed_test_recipient": "",
        }
    )
    return CampaignService(
        db,
        mailer=FakeMailer(),
        sleep_fn=lambda _: None,
        export_dir=tmp_path / "exports",
    )


def write_workbook(path: Path, rows: list[list[str]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    workbook.save(path)


def test_main_window_uses_russian_tabs_and_buttons(qapp: QApplication, tmp_path: Path) -> None:
    window = MainWindow(make_service(tmp_path))

    assert [window.tabs.tabText(index) for index in range(window.tabs.count())] == [
        "Рассылка",
        "Получатели",
        "Шаблоны",
        "Аккаунты и настройки",
        "Журнал",
    ]
    assert window.campaign_view.import_button.text() == "Загрузить Excel/CSV"
    assert window.campaign_view.send_button.text() == "Отправить подтвержденные"
    assert window.contacts_view.blacklist_button.text() == "В черный список"


def test_statuses_are_displayed_in_russian(qapp: QApplication, tmp_path: Path) -> None:
    service = make_service(tmp_path)
    campaign_id = service.default_campaign_id()
    contact_id = service.db.add_contact(
        {
            "campaign_id": campaign_id,
            "email": "status@example.com",
            "status": "dry_run_sent",
            "generated_message": "Текст",
        }
    )
    window = MainWindow(service)

    for row in range(window.contacts_view.table.rowCount()):
        if window.contacts_view.table.item(row, 0).text() == str(contact_id):
            assert window.contacts_view.table.item(row, 7).text() == status_to_ru("dry_run_sent")
            break
    else:
        raise AssertionError("contact row not rendered")


def test_import_russian_excel_with_subject_and_message(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    campaign_id = service.default_campaign_id()
    path = tmp_path / "contacts.xlsx"
    write_workbook(
        path,
        [
            ["Email", "Тема письма", "Сообщение", "Имя", "Компания", "Заметка"],
            ["a@example.com", "Сотрудничество", "Привет, хочу предложить...", "Анна", "Acme", "Ниша"],
        ],
    )

    result = service.import_file(path, campaign_id)
    contact = service.contacts(campaign_id)[0]

    assert result.imported_count == 1
    assert contact["subject"] == "Сотрудничество"
    assert contact["base_message"] == "Привет, хочу предложить..."
    assert contact["generated_message"] == "Привет, хочу предложить..."


def test_import_excel_with_email_and_message_only(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    campaign_id = service.default_campaign_id()
    path = tmp_path / "contacts.xlsx"
    write_workbook(path, [["Email", "Сообщение"], ["b@example.com", "Индивидуальный текст"]])

    result = service.import_file(path, campaign_id)
    contact = service.contacts(campaign_id)[0]

    assert result.imported_count == 1
    assert contact["email"] == "b@example.com"
    assert contact["generated_message"] == "Индивидуальный текст"


def test_individual_message_is_not_overwritten_by_template(tmp_path: Path) -> None:
    service = make_service(tmp_path)
    campaign_id = service.default_campaign_id()
    service.save_template("Шаблон темы", "ШАБЛОН НЕ ДОЛЖЕН ПЕРЕТЕРЕТЬ")
    service.add_contact_row(
        campaign_id,
        {
            "email": "personal@example.com",
            "subject": "Личная тема",
            "generated_message": "Личное сообщение",
        },
    )

    service.enqueue_generate_messages(campaign_id)
    QueueJobProcessor(service).process_available()
    contact = service.contacts(campaign_id)[0]

    assert contact["generated_message"] == "Личное сообщение"
    assert contact["subject"] == "Личная тема"
    assert contact["status"] == "pending_review"


def test_clipboard_paste_two_column_email_message_works(
    qapp: QApplication,
    tmp_path: Path,
    quiet_message_boxes: list[tuple[str, str]],
) -> None:
    service = make_service(tmp_path)
    window = MainWindow(service)
    QApplication.clipboard().setText("clip@example.com\tСообщение из буфера")

    window.contacts_view.paste_from_clipboard()

    contact = service.contacts(service.default_campaign_id())[0]
    assert contact["email"] == "clip@example.com"
    assert contact["generated_message"] == "Сообщение из буфера"
    assert quiet_message_boxes[-1][0] == "Вставка завершена"


def test_manual_add_and_delete_row(
    qapp: QApplication,
    tmp_path: Path,
    quiet_message_boxes: list[tuple[str, str]],
) -> None:
    service = make_service(tmp_path)
    window = MainWindow(service)

    window.contacts_view.add_empty_row()
    row = window.contacts_view.table.rowCount() - 1
    window.contacts_view.table.item(row, 1).setText("manual@example.com")
    window.contacts_view.table.item(row, 2).setText("Ручная тема")
    window.contacts_view.table.item(row, 3).setText("Ручное сообщение")
    window.contacts_view.save_changes()

    contacts = service.contacts(service.default_campaign_id())
    assert len(contacts) == 1
    assert contacts[0]["email"] == "manual@example.com"

    window.refresh_all()
    window.contacts_view.table.selectRow(0)
    window.contacts_view.delete_selected_rows()

    assert service.contacts(service.default_campaign_id()) == []
    assert quiet_message_boxes[-1][0] == "Удалить выбранные"


def test_example_excel_is_created(tmp_path: Path) -> None:
    service = make_service(tmp_path)

    path = service.export_example_contacts_template()
    workbook = load_workbook(path, read_only=True)
    sheet = workbook["Получатели"]

    assert path.name == "example_contacts_template.xlsx"
    assert sheet["A1"].value == "Email"
    assert sheet["C1"].value == "Сообщение"
    assert sheet["A2"].value.endswith("@example.com")


def test_settings_show_test_mode_instead_of_technical_dry_run(
    qapp: QApplication,
    tmp_path: Path,
) -> None:
    view = SettingsView(make_service(tmp_path), refresh_callback=lambda: None)
    view.refresh()

    assert view.send_mode.currentText() == "Тестовый режим"
    assert view.send_mode.currentData() == "dry_run"


def test_gmail_check_message_is_user_safe_in_russian(
    qapp: QApplication,
    tmp_path: Path,
    quiet_message_boxes: list[tuple[str, str]],
) -> None:
    view = SettingsView(make_service(tmp_path), refresh_callback=lambda: None)

    view._show_check_result(
        ConnectionCheckResult(
            False,
            "Gmail app password is missing. Save it in Email settings or configure .env for development.",
        )
    )

    assert "Gmail не подключен" in quiet_message_boxes[-1][1]
    assert "App Password" in quiet_message_boxes[-1][1]
    assert user_safe_error("Sender email is missing in Settings.").startswith("Не указан email")
