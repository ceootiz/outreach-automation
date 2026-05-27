from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .db import Database


EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

HEADER_ALIASES: dict[str, set[str]] = {
    "email": {"email", "mail", "почта", "e-mail", "адрес", "email адрес", "почтовый адрес"},
    "channel": {"channel", "канал", "канал рассылки", "platform", "платформа"},
    "handle": {"handle", "username", "user", "ник", "профиль username", "профиль / username", "x username", "instagram username", "telegram username"},
    "profile_url": {"profile url", "profile_url", "url профиля", "url profile", "ссылка профиля", "профиль url"},
    "external_id": {"external id", "external_id", "chat id", "user id", "id / chat", "id", "chat_id", "id chat"},
    "name": {"name", "имя"},
    "company": {"company", "компания"},
    "website": {"website", "site", "url", "сайт"},
    "social_profile": {"social", "social profile", "profile", "соцсеть", "профиль", "соцсеть / профиль"},
    "topic": {"topic", "niche", "note", "заметка", "ниша", "тема/ниша"},
    "subject": {"subject", "тема", "тема письма", "заголовок"},
    "base_message": {
        "message",
        "text",
        "body",
        "base_message",
        "сообщение",
        "текст",
        "письмо",
        "текст письма",
    },
}


@dataclass(slots=True)
class ImportResult:
    imported_count: int = 0
    skipped_count: int = 0
    errors: list[str] = field(default_factory=list)


def normalize_header(value: object) -> str:
    text = "" if value is None else str(value)
    return " ".join(text.strip().lower().replace("_", " ").split())


def map_headers(headers: list[object]) -> dict[str, int]:
    normalized = [normalize_header(header) for header in headers]
    mapping: dict[str, int] = {}
    for canonical, aliases in HEADER_ALIASES.items():
        normalized_aliases = {normalize_header(alias) for alias in aliases}
        for index, header in enumerate(normalized):
            if header in normalized_aliases:
                mapping[canonical] = index
                break
    return mapping


def is_valid_email(email: str) -> bool:
    return bool(EMAIL_RE.match(email.strip()))


class ContactImporter:
    def __init__(self, db: Database):
        self.db = db

    def import_file(self, file_path: str | Path, campaign_id: int) -> ImportResult:
        path = Path(file_path)
        rows = self._read_rows(path)
        result = ImportResult()
        seen_emails: set[str] = set()

        for row_number, row in rows:
            email = row.get("email", "").strip().lower()
            if not email or not is_valid_email(email):
                result.skipped_count += 1
                message = "Invalid or missing email"
                result.errors.append(f"Row {row_number}: {message}")
                self.db.log_import_error(campaign_id, path.name, row_number, email, message)
                continue

            if email in seen_emails or self.db.get_contact_by_email(campaign_id, email):
                result.skipped_count += 1
                message = "Duplicate email inside campaign"
                result.errors.append(f"Row {row_number}: {message} ({email})")
                self.db.log_import_error(campaign_id, path.name, row_number, email, message)
                continue

            seen_emails.add(email)
            contact = {
                "campaign_id": campaign_id,
                "channel": row.get("channel", "email") or "email",
                "email": email,
                "handle": row.get("handle", ""),
                "profile_url": row.get("profile_url", ""),
                "external_id": row.get("external_id", ""),
                "name": row.get("name", ""),
                "company": row.get("company", ""),
                "website": row.get("website", ""),
                "social_profile": row.get("social_profile", ""),
                "topic": row.get("topic", ""),
                "subject": row.get("subject", ""),
                "base_message": row.get("base_message", ""),
                "generated_message": row.get("base_message", ""),
                "last_error": (
                    ""
                    if row.get("base_message", "").strip()
                    else "Нет сообщения. Можно заполнить вручную или использовать шаблон."
                ),
                "status": "new",
            }
            try:
                self.db.add_contact(contact)
                result.imported_count += 1
            except Exception as exc:
                result.skipped_count += 1
                message = f"Failed to import row: {exc}"
                result.errors.append(f"Row {row_number}: {message}")
                self.db.log_import_error(campaign_id, path.name, row_number, email, message)

        return result

    def _read_rows(self, path: Path) -> list[tuple[int, dict[str, str]]]:
        suffix = path.suffix.lower()
        if suffix in {".xlsx", ".xlsm"}:
            return self._read_excel(path)
        if suffix == ".csv":
            return self._read_csv(path)
        raise ValueError(f"Unsupported file type: {suffix}")

    def _read_excel(self, path: Path) -> list[tuple[int, dict[str, str]]]:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator, None)
        if not headers:
            return []
        mapping = map_headers(list(headers))
        if "email" not in mapping:
            raise ValueError("Required column 'email' was not found")

        rows: list[tuple[int, dict[str, str]]] = []
        for excel_row_number, values in enumerate(iterator, start=2):
            values_list = list(values or [])
            if not any(value not in (None, "") for value in values_list):
                continue
            rows.append((excel_row_number, self._extract_row(mapping, values_list)))
        return rows

    def _read_csv(self, path: Path) -> list[tuple[int, dict[str, str]]]:
        with path.open("r", encoding="utf-8-sig", newline="") as file:
            reader = csv.reader(file)
            headers = next(reader, None)
            if not headers:
                return []
            mapping = map_headers(headers)
            if "email" not in mapping:
                raise ValueError("Required column 'email' was not found")

            rows: list[tuple[int, dict[str, str]]] = []
            for csv_row_number, values in enumerate(reader, start=2):
                if not any(value.strip() for value in values if value is not None):
                    continue
                rows.append((csv_row_number, self._extract_row(mapping, values)))
        return rows

    @staticmethod
    def _extract_row(mapping: dict[str, int], values: list[Any]) -> dict[str, str]:
        row: dict[str, str] = {}
        for field_name, index in mapping.items():
            value = values[index] if index < len(values) else ""
            row[field_name] = "" if value is None else str(value).strip()
        return row
