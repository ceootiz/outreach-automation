from __future__ import annotations

import csv
import time
import json
import hashlib
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any, Callable

from openpyxl import Workbook

from .ai import AIConnectionCheckResult, AIService, DualBrainResult, DualBrainService
from .ai.brain_config import BrainConfig, normalize_brain_provider, normalize_generation_mode
from .ai.research_schema import RecipientBrief, parse_recipient_brief
from .ai_provider import BaseAIProvider, PlaceholderAIProvider
from .blacklist import BlacklistService
from .channels import (
    BaseChannel,
    ChannelSafetyError,
    build_connector_slots,
    channel_options,
    get_channel,
    list_channel_readiness,
    list_channels,
    recommend_channel_for_lead,
)
from .channels.execution import (
    DRY_RUN,
    MANUAL_ASSIST,
    OFFICIAL_API,
    EXECUTION_MODE_LABELS,
    ChannelCapability,
    ChannelExecutionEngine,
    ChannelExecutionResult,
    ManualAssistAction,
    get_channel_capability,
)
from .channels.telegram_channel import TelegramChannel
from .config import EXPORTS_DIR, as_bool, as_int, get_gmail_app_password, get_telegram_bot_token
from .credential_store import (
    CredentialStatus,
    delete_gmail_app_password,
    get_gmail_credential_status,
    get_telegram_credential_status,
    has_ai_api_key,
    has_ai_brain_api_key,
    load_ai_api_key,
    load_ai_brain_api_key,
    load_telegram_bot_token,
    save_ai_api_key,
    save_ai_brain_api_key,
    save_gmail_app_password,
    save_telegram_bot_token,
)
from .db import Database
from .db_safety import backup_database
from .excel_importer import ContactImporter, ImportResult, is_valid_email
from .enrichment import EnrichmentResult, EnrichmentService
from .gmail_profile_service import GmailProfileService
from .inbox import EmailSyncClient, InboxIngestionService, InboxSyncResult, TelegramSyncClient
from .intelligence import (
    AIQualityService,
    AnalyticsService,
    ConversationAI,
    ConversationService,
    ContactTimelineService,
    FollowUpService,
    MetricsService,
    ReplyService,
)
from .logger_setup import get_logger, redact_secret
from .mailer import ConnectionCheckResult, GmailSMTPMailer, MailerConfigError, MailerError
from .operator import (
    HIGH_VOLUME_MODE,
    HOTKEY_ACTIONS,
    OPERATOR_MODES,
    PRECISION_MODE,
    LeadPrioritizer,
    OperatorSessionManager,
    ReviewQueueBuilder,
    ReviewQueueFilters,
    normalize_operator_mode,
)
from .performance import CacheManager, LatencyMetrics, SessionCache, chunked, recommended_batch_size, virtual_page
from .presets import CampaignPreset, get_preset, load_presets
from .queue_service import QueueResult, QueueService
from .rate_limiter import RateLimiter
from .search import GlobalSearchService


@dataclass(slots=True)
class SendSummary:
    sent: int = 0
    dry_run_sent: int = 0
    failed: int = 0
    blacklisted: int = 0
    blocked_by_limit: bool = False
    blocked_by_guardrail: bool = False
    errors: list[str] = field(default_factory=list)


class CampaignService:
    def __init__(
        self,
        db: Database,
        mailer: GmailSMTPMailer | None = None,
        ai_provider: BaseAIProvider | None = None,
        ai_draft_service: AIService | None = None,
        dual_brain_service: DualBrainService | None = None,
        enrichment_service: EnrichmentService | None = None,
        sleep_fn: Callable[[float], None] = time.sleep,
        export_dir: Path | None = None,
        backup_dir: Path | None = None,
    ):
        self.db = db
        self.mailer = mailer or GmailSMTPMailer()
        self.ai_provider = ai_provider or PlaceholderAIProvider()
        self.ai_draft_service = ai_draft_service or AIService()
        self.dual_brain_service = dual_brain_service or DualBrainService()
        self.enrichment = enrichment_service or EnrichmentService(db)
        self.importer = ContactImporter(db)
        self.blacklist = BlacklistService(db)
        self.rate_limiter = RateLimiter(db)
        self.queue = QueueService(db)
        self.execution = ChannelExecutionEngine()
        self.lead_prioritizer = LeadPrioritizer()
        self.review_queue = ReviewQueueBuilder(self.lead_prioritizer)
        self.operator_sessions = OperatorSessionManager(db)
        self.gmail_profiles = GmailProfileService(db, self.mailer)
        self.timeline = ContactTimelineService(db)
        self.replies = ReplyService(db, self.timeline)
        self.ai_quality = AIQualityService(db, self.timeline)
        self.metrics = MetricsService(db)
        self.analytics = AnalyticsService(db, self.metrics, self.timeline)
        self.followups = FollowUpService(db, self.timeline)
        self.conversation_ai = ConversationAI()
        self.conversations = ConversationService(db, self.timeline, self.replies, self.conversation_ai)
        self.inbox = InboxIngestionService(db, self.conversations, self.queue)
        self.search_service = GlobalSearchService(db)
        self.cache = CacheManager(default_ttl_seconds=45, max_items=512)
        self.session_cache = SessionCache(db)
        self.latency_metrics = LatencyMetrics()
        self.sleep_fn = sleep_fn
        self.export_dir = export_dir or EXPORTS_DIR
        self.backup_dir = backup_dir
        self.logger = get_logger()
        self.gmail_profiles.migrate_legacy_credentials_if_needed()
        recovered = self.queue.recover_interrupted_jobs()
        if recovered:
            self.db.log_send(
                None,
                None,
                "queue_recovery",
                "ok",
                f"Recovered interrupted jobs after restart: {recovered}.",
            )

    def backup_before_risky_operation(self, reason: str) -> None:
        if self.backup_dir:
            backup_database(self.db.path, self.backup_dir, reason=reason)

    def campaigns(self) -> list[dict[str, Any]]:
        return self.db.get_campaigns()

    def default_campaign_id(self) -> int:
        return self.db.get_default_campaign_id()

    def campaign_presets(self) -> list[CampaignPreset]:
        return load_presets()

    def campaign_preset(self, preset_id: str) -> CampaignPreset:
        return get_preset(preset_id)

    def create_campaign_from_preset(self, preset_id: str, name: str | None = None) -> int:
        preset = self.campaign_preset(preset_id)
        campaign_name = (name or preset.name).strip() or preset.name
        campaign_id = self.db.create_campaign(campaign_name)
        self.apply_campaign_preset(preset.preset_id, campaign_id=campaign_id)
        self.db.log_send(
            None,
            campaign_id,
            "campaign_created",
            "ok",
            f"preset={preset.preset_id}; channel={preset.recommended_channel}; autosend=false",
            channel=preset.recommended_channel,
        )
        return campaign_id

    def apply_campaign_preset(self, preset_id: str, campaign_id: int | None = None) -> dict[str, Any]:
        preset = self.campaign_preset(preset_id)
        channel_id = self.set_active_channel(preset.recommended_channel)
        execution_mode = self.set_execution_mode(channel_id, preset.recommended_execution_mode)
        settings: dict[str, str] = {
            "active_campaign_preset": preset.preset_id,
            "active_channel": channel_id,
            "work_mode": "ai_assist" if preset.ai_enabled else "manual",
            "ai_tone": preset.ai_tone,
            "send_mode": "dry_run",
            "operator_quick_start": "",
        }
        if preset.recommended_limits.get("daily_send_limit"):
            settings["daily_send_limit"] = str(preset.recommended_limits["daily_send_limit"])
        if preset.recommended_limits.get("follow_up_days"):
            settings["follow_up_delay_days"] = str(preset.recommended_limits["follow_up_days"])
        self.save_settings(settings)
        target_campaign_id = campaign_id or self.default_campaign_id()
        self.db.log_send(
            None,
            target_campaign_id,
            "apply_campaign_preset",
            "ok",
            f"preset={preset.preset_id}; channel={channel_id}; execution={execution_mode}; autosend=false",
            channel=channel_id,
        )
        return {
            "preset": preset.to_dict(),
            "campaign_id": target_campaign_id,
            "channel": channel_id,
            "execution_mode": execution_mode,
            "autosend": False,
        }

    def archive_campaign(self, campaign_id: int) -> None:
        self.backup_before_risky_operation("campaign_archive")
        self.db.execute(
            "UPDATE campaigns SET status = 'archived' WHERE id = ?",
            (campaign_id,),
        )
        self.db.log_send(None, campaign_id, "campaign_archive", "archived", "Campaign archived by operator.")

    def duplicate_campaign(self, campaign_id: int, name: str | None = None) -> int:
        self.backup_before_risky_operation("campaign_duplicate")
        source = self.db.fetch_one("SELECT * FROM campaigns WHERE id = ?", (campaign_id,))
        if not source:
            raise ValueError("Campaign not found.")
        duplicate_name = (name or f"{source.get('name') or 'Campaign'} copy").strip()
        new_campaign_id = self.db.create_campaign(duplicate_name)
        contacts = self.contacts(campaign_id)
        for contact in contacts:
            new_contact_id = self.db.add_contact(
                {
                    "campaign_id": new_campaign_id,
                    "channel": contact.get("channel") or "email",
                    "email": contact.get("email") or "",
                    "handle": contact.get("handle") or "",
                    "profile_url": contact.get("profile_url") or "",
                    "external_id": contact.get("external_id") or "",
                    "name": contact.get("name") or "",
                    "company": contact.get("company") or "",
                    "topic": contact.get("topic") or "",
                    "website": contact.get("website") or "",
                    "social_profile": contact.get("social_profile") or "",
                    "subject": contact.get("subject") or "",
                    "base_message": contact.get("base_message") or "",
                    "generated_message": contact.get("generated_message") or "",
                    "ai_generated": contact.get("ai_generated") or 0,
                    "ai_confidence": contact.get("ai_confidence"),
                    "ai_notes": contact.get("ai_notes") or "",
                    "ai_warnings": contact.get("ai_warnings") or "",
                    "research_brief_json": contact.get("research_brief_json") or "",
                    "research_confidence": contact.get("research_confidence"),
                    "research_warnings": contact.get("research_warnings") or "",
                    "research_status": contact.get("research_status") or "",
                    "enrichment_result_json": contact.get("enrichment_result_json") or "",
                    "enrichment_status": contact.get("enrichment_status") or "not_checked",
                    "enrichment_source_urls": contact.get("enrichment_source_urls") or "",
                    "enrichment_warnings": contact.get("enrichment_warnings") or "",
                    "enrichment_confidence": contact.get("enrichment_confidence"),
                    "status": "new",
                    "last_error": "Duplicated campaign draft. Review before sending.",
                }
            )
            self.timeline.record_event(
                new_contact_id,
                "contact_created",
                title="Контакт скопирован",
                details=f"Copied from campaign {campaign_id}. Human review required.",
                metadata={"source_campaign_id": campaign_id},
            )
            self.conversations.get_or_create_thread(new_contact_id)
        self.db.log_send(
            None,
            new_campaign_id,
            "campaign_duplicate",
            "ok",
            f"source_campaign_id={campaign_id}; copied_contacts={len(contacts)}; autosend=false",
        )
        return new_campaign_id

    def quick_start_campaign(self, quick_start_id: str) -> dict[str, Any]:
        mapping = {
            "quick_email_outreach": ("b2b_email_outreach", "Quick Email Outreach"),
            "quick_telegram_campaign": ("telegram_outreach", "Quick Telegram Campaign"),
            "quick_ai_draft_generation": ("partnership_intro", "Quick AI Draft Generation"),
        }
        key = (quick_start_id or "").strip().lower()
        if key not in mapping:
            raise ValueError("Unknown quick start mode.")
        preset_id, name = mapping[key]
        campaign_id = self.create_campaign_from_preset(preset_id, name=name)
        extra_settings = {"operator_quick_start": key, "send_mode": "dry_run"}
        if key == "quick_ai_draft_generation":
            extra_settings.update({"work_mode": "ai_assist", "ai_generation_mode": "dual_brain"})
        self.save_settings(extra_settings)
        return {
            "campaign_id": campaign_id,
            "quick_start_id": key,
            "preset": self.campaign_preset(preset_id).to_dict(),
            "autosend": False,
        }

    def campaign_validation(self, campaign_id: int, preset_id: str | None = None) -> dict[str, Any]:
        settings = self.settings()
        channel_id = self.active_channel_id()
        execution_mode = self.execution_mode(channel_id)
        preset = self.campaign_preset(preset_id or settings.get("active_campaign_preset", "b2b_email_outreach"))
        stats = self.stats(campaign_id)
        checks: list[dict[str, str]] = []

        def add(check_id: str, label: str, status: str, message: str) -> None:
            checks.append({"id": check_id, "label": label, "status": status, "message": message})

        if stats.get("total", 0) > 0:
            add("recipients", "Получатели импортированы", "ok", f"Получателей: {stats.get('total', 0)}.")
        else:
            add("recipients", "Получатели импортированы", "error", "Добавьте получателей перед запуском кампании.")

        if channel_id == "email":
            sender = self.active_sender_email()
            sender_status = self.gmail_credential_status(sender)
            if sender and sender_status.has_password:
                add("sender", "Активный отправитель", "ok", f"Email sender: {sender}.")
            else:
                add("sender", "Активный отправитель", "error", "Выберите Gmail-профиль и сохраните App Password.")
        elif channel_id == "telegram":
            if self.telegram_credential_status().has_password:
                add("telegram", "Telegram Bot подключен", "ok", "Bot Token сохранен безопасно.")
            else:
                add("telegram", "Telegram Bot подключен", "warning", "Для live Telegram нужен Bot Token; dry-run доступен.")
        else:
            capability = self.channel_capability(channel_id)
            if execution_mode == MANUAL_ASSIST:
                add("execution", "Manual Assist выбран", "ok", capability.limitation_text)
            else:
                add("execution", "Manual Assist выбран", "warning", f"{capability.display_name}: лучше использовать Manual Assist.")

        ai_needed = preset.ai_enabled or settings.get("work_mode") == "ai_assist"
        if ai_needed:
            generation_mode = normalize_generation_mode(settings.get("ai_generation_mode", "simple"))
            if generation_mode == "dual_brain":
                research_ok = self.ai_brain_credential_status("research", settings.get("ai_research_provider", "openai")).has_password
                writer_ok = self.ai_brain_credential_status("writer", settings.get("ai_writer_provider", "openai")).has_password
                if research_ok and writer_ok:
                    add("ai", "AI настроен", "ok", "Research Brain и Writer Brain готовы.")
                else:
                    add("ai", "AI настроен", "warning", "Добавьте API key для Research Brain и Writer Brain.")
            elif self.ai_credential_status(settings.get("ai_provider", "openai")).has_password:
                add("ai", "AI настроен", "ok", "AI Assist готов.")
            else:
                add("ai", "AI настроен", "warning", "AI drafts доступны после сохранения API key.")
        else:
            add("ai", "AI настроен", "ok", "Ручной режим: AI не обязателен.")

        if as_bool(settings.get("safe_mode"), True):
            add("safe_mode", "Safe mode включен", "ok", "Live send защищен дополнительными guardrails.")
        else:
            add("safe_mode", "Safe mode включен", "warning", "Рекомендуется включить safe mode перед live send.")

        daily_limit = as_int(settings.get("daily_send_limit"), 25)
        if daily_limit > 0:
            add("daily_limit", "Daily limit задан", "ok", f"Лимит: {daily_limit}.")
        else:
            add("daily_limit", "Daily limit задан", "error", "Daily limit должен быть больше 0.")

        follow_up_days = as_int(settings.get("follow_up_delay_days"), 0)
        if follow_up_days > 0:
            add("follow_up", "Follow-up стратегия", "ok", f"Напоминание через {follow_up_days} дн.")
        else:
            add("follow_up", "Follow-up стратегия", "warning", "Follow-up не настроен.")

        return {
            "campaign_id": campaign_id,
            "preset": preset.to_dict(),
            "channel": channel_id,
            "execution_mode": execution_mode,
            "checks": checks,
            "warnings": [row["message"] for row in checks if row["status"] == "warning"],
            "errors": [row["message"] for row in checks if row["status"] == "error"],
        }

    def campaign_health(self, campaign_id: int, preset_id: str | None = None) -> dict[str, Any]:
        validation = self.campaign_validation(campaign_id, preset_id=preset_id)
        score = 100
        for check in validation["checks"]:
            if check["status"] == "error":
                score -= 25
            elif check["status"] == "warning":
                score -= 10
        score = max(0, min(100, score))
        if score >= 80:
            label = "Healthy"
            status = "healthy"
        elif score >= 55:
            label = "Needs attention"
            status = "needs_attention"
        else:
            label = "Risky"
            status = "risky"
        return {
            "campaign_id": campaign_id,
            "score": score,
            "label": label,
            "status": status,
            "validation": validation,
        }

    def smart_warnings(
        self,
        campaign_id: int,
        *,
        channel: str | None = None,
        execution_mode: str | None = None,
        send_mode: str | None = None,
    ) -> list[str]:
        channel_id = (channel or self.active_channel_id()).strip().lower()
        mode = execution_mode or self.execution_mode(channel_id)
        current_send_mode = (send_mode or self.settings().get("send_mode") or "dry_run").strip().lower()
        capability = self.channel_capability(channel_id)
        stats = self.stats(campaign_id)
        warnings: list[str] = []
        if stats.get("total", 0) == 0:
            warnings.append("Добавьте получателей перед запуском кампании.")
        if channel_id in {"instagram", "x", "tiktok", "vk"} and mode != MANUAL_ASSIST:
            warnings.append(f"{capability.display_name}: автоматическая отправка отключена. Используйте Manual Assist.")
        if channel_id == "telegram":
            missing_chat = self.db.fetch_one(
                """
                SELECT COUNT(*) AS count
                FROM contacts
                WHERE campaign_id = ? AND channel = 'telegram' AND COALESCE(external_id, '') = ''
                """,
                (campaign_id,),
            )
            if int(missing_chat["count"] if missing_chat else 0) > 0:
                warnings.append("У части Telegram получателей нет chat_id. Live send будет заблокирован.")
        if current_send_mode == "live" and stats.get("approved", 0) > 25 and stats.get("dry_run_sent", 0) == 0:
            warnings.append(f"Вы собираетесь отправить {stats.get('approved', 0)} сообщений без dry-run.")
        low_ai = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM contacts
            WHERE campaign_id = ?
              AND ai_generated = 1
              AND (ai_confidence IS NULL OR ai_confidence < 0.45 OR COALESCE(ai_warnings, '') <> '')
            """,
            (campaign_id,),
        )
        low_ai_count = int(low_ai["count"] if low_ai else 0)
        if low_ai_count:
            warnings.append(f"AI confidence low или есть warnings у {low_ai_count} контактов.")
        if current_send_mode == "live" and not as_bool(self.settings().get("safe_mode"), True):
            warnings.append("Safe mode выключен. Рекомендуется включить его перед live send.")
        return warnings

    def operator_dashboard(self, campaign_id: int | None = None) -> dict[str, Any]:
        active_campaigns = [
            row for row in self.campaigns()
            if str(row.get("status") or "active") != "archived"
        ]
        target_campaign_id = campaign_id or self.default_campaign_id()
        stats = self.stats(target_campaign_id)
        replies_waiting = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM replies
            WHERE campaign_id = ? AND reply_status IN ('interested', 'maybe_later', 'follow_up_needed', 'no_response')
            """,
            (target_campaign_id,),
        )
        followups_due = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM followups
            WHERE campaign_id = ? AND status IN ('scheduled', 'queued')
            """,
            (target_campaign_id,),
        )
        ai_warnings = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM contacts
            WHERE campaign_id = ? AND (COALESCE(ai_warnings, '') <> '' OR COALESCE(research_warnings, '') <> '')
            """,
            (target_campaign_id,),
        )
        warnings = self.smart_warnings(target_campaign_id)
        return {
            "campaign_id": target_campaign_id,
            "active_campaigns": len(active_campaigns),
            "drafts_pending_review": stats.get("pending_review", 0),
            "replies_waiting": int(replies_waiting["count"] if replies_waiting else 0),
            "followup_reminders": int(followups_due["count"] if followups_due else 0),
            "risk_alerts": len(warnings),
            "ai_warnings": int(ai_warnings["count"] if ai_warnings else 0),
            "health": self.campaign_health(target_campaign_id),
            "warnings": warnings,
        }

    def operator_mode_options(self) -> list[tuple[str, str]]:
        return [(key, label) for key, label in OPERATOR_MODES.items()]

    def operator_mode(self) -> str:
        return normalize_operator_mode(self.settings().get("operator_mode", PRECISION_MODE))

    def set_operator_mode(self, mode: str) -> str:
        normalized = normalize_operator_mode(mode)
        self.save_settings({"operator_mode": normalized})
        return normalized

    def operator_hotkeys(self) -> dict[str, dict[str, str]]:
        return {
            key: {
                "action_id": value.action_id,
                "label": value.label,
                "description": value.description,
            }
            for key, value in HOTKEY_ACTIONS.items()
        }

    def operator_priority_for_contact(self, contact_id: int) -> dict[str, Any]:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        return self.lead_prioritizer.score_contact(contact).to_dict()

    def operator_review_queue(
        self,
        campaign_id: int,
        filters: ReviewQueueFilters | dict[str, Any] | None = None,
    ) -> list[dict[str, Any]]:
        selected_filters = filters if isinstance(filters, ReviewQueueFilters) else ReviewQueueFilters.from_dict(filters)
        rows = self.contacts(campaign_id)
        return [item.to_dict() for item in self.review_queue.build(rows, selected_filters)]

    def operator_review_queue_page(
        self,
        campaign_id: int,
        filters: ReviewQueueFilters | dict[str, Any] | None = None,
        *,
        offset: int = 0,
        limit: int = 250,
    ) -> dict[str, Any]:
        queue = self.operator_review_queue(campaign_id, filters)
        page = virtual_page(len(queue), offset=offset, limit=limit)
        return {
            "items": queue[page.offset: page.offset + page.limit],
            "page": page.to_dict(),
            "render_budget": recommended_batch_size(len(queue)).to_dict(),
        }

    def start_outreach_session(
        self,
        campaign_id: int,
        *,
        mode: str | None = None,
        filters: ReviewQueueFilters | dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        selected_filters = filters if isinstance(filters, ReviewQueueFilters) else ReviewQueueFilters.from_dict(filters)
        queue_items = self.operator_review_queue(campaign_id, selected_filters)
        first_contact_id = (
            int(queue_items[0]["contact"]["id"])
            if queue_items
            else None
        )
        selected_mode = normalize_operator_mode(mode or self.operator_mode() or HIGH_VOLUME_MODE)
        self.set_operator_mode(selected_mode)
        state = self.operator_sessions.start_session(
            campaign_id,
            mode=selected_mode,
            filters=selected_filters,
            current_contact_id=first_contact_id,
        )
        self.db.log_send(
            None,
            campaign_id,
            "operator_session_start",
            "active",
            f"mode={selected_mode}; queue={len(queue_items)}; no autosend",
        )
        return self.operator_session_snapshot(state.id)

    def restore_outreach_session(self, campaign_id: int | None = None) -> dict[str, Any] | None:
        state = self.operator_sessions.restore_session(campaign_id=campaign_id)
        return self.operator_session_snapshot(state.id) if state else None

    def save_operator_ui_state(self, **state: Any) -> dict[str, Any]:
        return self.session_cache.save(**state)

    def restore_operator_ui_state(self) -> dict[str, Any]:
        return self.session_cache.load()

    def operator_session_snapshot(
        self,
        session_id: int,
        filters: ReviewQueueFilters | dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        state = self.operator_sessions.restore_session(session_id=session_id)
        if not state:
            raise ValueError("Operator session not found.")
        explicit_filters = filters is not None
        selected_filters = filters if isinstance(filters, ReviewQueueFilters) else (
            ReviewQueueFilters.from_dict(filters) if filters is not None else state.filters
        )
        queue_items = self.operator_review_queue(state.campaign_id, selected_filters)
        if queue_items:
            index = min(max(state.review_index, 0), len(queue_items) - 1)
            if state.current_contact_id:
                for item_index, item in enumerate(queue_items):
                    if int(item["contact"]["id"]) == state.current_contact_id:
                        index = item_index
                        break
            current_item = queue_items[index]
            current_contact_id = int(current_item["contact"]["id"])
            if current_contact_id != state.current_contact_id or index != state.review_index:
                self.operator_sessions.save_position(
                    state.id,
                    contact_id=current_contact_id,
                    review_index=index,
                    filters=selected_filters,
                    draft_variant=state.draft_variant,
                    unsaved_draft=state.unsaved_draft if current_contact_id == state.current_contact_id else "",
                )
                state = self.operator_sessions.restore_session(session_id=state.id) or state
        else:
            index = 0
            current_item = None
        if explicit_filters:
            self.operator_sessions.save_position(
                state.id,
                contact_id=state.current_contact_id,
                review_index=index,
                filters=selected_filters,
                draft_variant=state.draft_variant,
                unsaved_draft=state.unsaved_draft,
            )
            state = self.operator_sessions.restore_session(session_id=state.id) or state
        return {
            "session": state.to_dict(),
            "queue": queue_items,
            "current": current_item,
            "index": index,
            "total": len(queue_items),
            "metrics": self.operator_sessions.metrics(state.id).to_dict(),
            "hotkeys": self.operator_hotkeys(),
            "rate_limit": self.operator_rate_limit_state(state.campaign_id),
            "autosend": False,
        }

    def operator_move_session(
        self,
        session_id: int,
        *,
        direction: str = "next",
    ) -> dict[str, Any]:
        state = self.operator_sessions.restore_session(session_id=session_id)
        if not state:
            raise ValueError("Operator session not found.")
        filters = state.filters
        queue_items = self.operator_review_queue(state.campaign_id, filters)
        if not queue_items:
            return self.operator_session_snapshot(session_id)
        index = min(max(state.review_index, 0), len(queue_items) - 1)
        if state.current_contact_id:
            for item_index, item in enumerate(queue_items):
                if int(item["contact"]["id"]) == state.current_contact_id:
                    index = item_index
                    break
        next_index = index + (1 if direction == "next" else -1)
        next_index = max(0, min(next_index, len(queue_items) - 1))
        contact_id = int(queue_items[next_index]["contact"]["id"])
        self.operator_sessions.save_position(
            session_id,
            contact_id=contact_id,
            review_index=next_index,
            filters=filters,
            draft_variant="",
            unsaved_draft="",
        )
        self.operator_sessions.record_action(session_id, contact_id, f"{direction}_lead")
        state = self.operator_sessions.restore_session(session_id=session_id) or state
        return {
            "session": state.to_dict(),
            "queue": queue_items,
            "current": queue_items[next_index],
            "index": next_index,
            "total": len(queue_items),
            "metrics": self.operator_sessions.metrics(state.id).to_dict(),
            "hotkeys": self.operator_hotkeys(),
            "rate_limit": self.operator_rate_limit_state(state.campaign_id),
            "autosend": False,
        }

    def operator_skip_lead(self, session_id: int, contact_id: int) -> dict[str, Any]:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        self.operator_sessions.record_action(session_id, contact_id, "skip_lead")
        self.timeline.record_event(
            contact_id,
            "note_added",
            title="Lead skipped",
            details="Operator skipped this lead during Outreach Session. No message was sent.",
            metadata={"session_id": session_id},
        )
        return self.operator_move_session(session_id, direction="next")

    def operator_copy_message(self, session_id: int, contact_id: int) -> str:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        state = self.operator_sessions.restore_session(session_id=session_id)
        if state and state.current_contact_id == contact_id and state.unsaved_draft:
            contact = {**contact, "generated_message": state.unsaved_draft}
        action = self.execution.manual_assist_action(contact)
        text = action.copy_text or action.message
        self.operator_sessions.record_action(
            session_id,
            contact_id,
            "copy_message",
            {"channel": action.channel, "manual_required": True},
        )
        self.db.log_send(
            contact_id,
            int(contact["campaign_id"]),
            "operator_copy_message",
            "manual_required",
            "Message copied for operator. Nothing was sent.",
            channel=action.channel,
            platform_recipient=action.recipient,
        )
        return text

    def operator_open_profile(self, session_id: int, contact_id: int) -> ManualAssistAction:
        action = self.prepare_manual_assist(contact_id)
        self.operator_sessions.record_action(
            session_id,
            contact_id,
            "open_profile",
            {"profile_url": action.profile_url, "manual_required": True},
        )
        return action

    def operator_mark_manually_sent(self, session_id: int, contact_id: int) -> dict[str, Any]:
        state = self.operator_sessions.restore_session(session_id=session_id)
        if state and state.current_contact_id == contact_id and state.unsaved_draft:
            self.db.update_contact(contact_id, {"generated_message": state.unsaved_draft})
        count = self.mark_manual_assist_sent([contact_id])
        self.operator_sessions.record_action(
            session_id,
            contact_id,
            "mark_manually_sent",
            {"count": count, "manual_required": True},
        )
        return self.operator_move_session(session_id, direction="next")

    def operator_approve_draft(self, session_id: int, contact_id: int) -> None:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        old_status = str(contact.get("status") or "")
        self.db.update_contact(contact_id, {"status": "approved", "last_error": ""})
        self.timeline.record_status_change(contact_id, old_status, "approved")
        self.operator_sessions.record_action(session_id, contact_id, "approve_draft")

    def operator_regenerate_draft(self, session_id: int, contact_id: int) -> QueueResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return QueueResult(False, error="Contact not found.")
        result = self.enqueue_ai_generate_drafts(
            int(contact["campaign_id"]),
            contact_ids=[contact_id],
            confirm_ai_cost=True,
            channel=str(contact.get("channel") or "email"),
        )
        self.operator_sessions.record_action(
            session_id,
            contact_id,
            "regenerate_draft",
            {"queued": result.count, "ok": result.ok, "autosend": False},
        )
        return result

    def operator_schedule_followup(self, session_id: int, contact_id: int) -> int:
        followup_id = self.schedule_followup(contact_id)
        self.operator_sessions.record_action(session_id, contact_id, "follow_up", {"followup_id": followup_id})
        return followup_id

    def operator_change_lead_status(self, session_id: int, contact_id: int, lead_status: str) -> str:
        status = self.change_lead_status(contact_id, lead_status)
        self.operator_sessions.record_action(session_id, contact_id, "change_lead_status", {"lead_status": status})
        return status

    def operator_select_variant(self, session_id: int, contact_id: int, variant: str) -> str:
        variants = self.operator_fast_variants(contact_id)
        key = variant if variant in variants else "short"
        text = variants[key]
        self.db.update_contact(contact_id, {"generated_message": text})
        state = self.operator_sessions.restore_session(session_id=session_id)
        self.operator_sessions.save_position(
            session_id,
            contact_id=contact_id,
            review_index=state.review_index if state else 0,
            draft_variant=key,
            unsaved_draft=text,
        )
        self.operator_sessions.record_action(session_id, contact_id, f"choose_variant_{key}")
        return text

    def operator_save_draft_state(self, session_id: int, contact_id: int, draft_text: str) -> None:
        state = self.operator_sessions.restore_session(session_id=session_id)
        if not state:
            raise ValueError("Operator session not found.")
        self.operator_sessions.save_position(
            session_id,
            contact_id=contact_id,
            review_index=state.review_index,
            filters=state.filters,
            draft_variant=state.draft_variant,
            unsaved_draft=draft_text,
        )

    def operator_fast_variants(self, contact_id: int) -> dict[str, str]:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        base = str(contact.get("generated_message") or contact.get("base_message") or "").strip()
        if not base:
            base = "Здравствуйте! Хочу аккуратно предложить сотрудничество. Если актуально, готов отправить детали."
        short = base if len(base) <= 320 else base[:317].rstrip() + "..."
        friendly = base
        if not friendly.lower().startswith(("привет", "здравствуйте", "добрый")):
            friendly = f"Здравствуйте! {friendly}"
        direct = base
        if not direct.lower().startswith(("коротко", "здравствуйте", "добрый")):
            direct = f"Коротко: {direct}"
        return {"short": short, "friendly": friendly, "direct": direct}

    def operator_save_ai_feedback(
        self,
        session_id: int | None,
        contact_id: int,
        rating: str,
        note: str = "",
    ) -> int:
        feedback_id = self.operator_sessions.save_feedback(
            contact_id=contact_id,
            rating=rating,
            note=note,
            session_id=session_id,
        )
        if session_id:
            self.operator_sessions.record_action(session_id, contact_id, "ai_feedback", {"rating": rating})
        return feedback_id

    def operator_rate_limit_state(self, campaign_id: int, channel: str | None = None) -> dict[str, Any]:
        settings = self.settings()
        limit_state = self.rate_limiter.state(settings.get("daily_send_limit", "25"))
        active_channel = (channel or self.active_channel_id()).strip().lower()
        warnings: list[str] = []
        if not limit_state.allowed:
            warnings.append("Daily limit reached. Session paused; no bypass attempts.")
        if active_channel in {"instagram", "x", "tiktok", "vk"}:
            warnings.append("Manual Assist only. The operator must send manually.")
        if not as_bool(settings.get("safe_mode"), True):
            warnings.append("Safe mode is off. Enable it before live execution.")
        return {
            "campaign_id": campaign_id,
            "channel": active_channel,
            "daily_limit": limit_state.daily_limit,
            "sent_today": limit_state.sent_today,
            "remaining": limit_state.remaining,
            "allowed": limit_state.allowed,
            "paused": not limit_state.allowed,
            "warnings": warnings,
        }

    def operator_analytics_summary(self, campaign_id: int) -> dict[str, Any]:
        events = self.db.fetch_all(
            """
            SELECT ose.action, ose.created_at, c.channel
            FROM operator_session_events ose
            LEFT JOIN contacts c ON c.id = ose.contact_id
            LEFT JOIN operator_sessions os ON os.id = ose.session_id
            WHERE os.campaign_id = ?
            ORDER BY ose.created_at ASC, ose.id ASC
            """,
            (campaign_id,),
        )
        reviewed = len([row for row in events if row["action"] in {"approve_draft", "copy_message", "mark_manually_sent"}])
        copied = len([row for row in events if row["action"] == "copy_message"])
        manually_sent = len([row for row in events if row["action"] == "mark_manually_sent"])
        followups = len([row for row in events if row["action"] in {"follow_up", "followup_done"}])
        accepted = len([row for row in events if row["action"] in {"approve_draft", "mark_manually_sent"}])
        by_channel: dict[str, int] = {}
        for row in events:
            channel = str(row.get("channel") or "unknown")
            by_channel[channel] = by_channel.get(channel, 0) + 1
        return {
            "reviewed": reviewed,
            "copied": copied,
            "manually_sent": manually_sent,
            "followups": followups,
            "ai_acceptance_rate": round(accepted / reviewed, 2) if reviewed else 0.0,
            "throughput_per_hour": reviewed,
            "channel_activity": by_channel,
        }

    def command_palette_items(self, query: str = "", campaign_id: int | None = None, limit: int = 30) -> list[dict[str, Any]]:
        needle = (query or "").strip()
        target_campaign = campaign_id or self.default_campaign_id()
        items: list[dict[str, Any]] = [
            {"id": "open:campaigns", "title": "Open Campaigns", "subtitle": "Кампании и checklist", "kind": "navigation"},
            {"id": "open:outreach_session", "title": "Open Outreach Session", "subtitle": "High Volume review", "kind": "navigation"},
            {"id": "open:inbox", "title": "Open Inbox", "subtitle": "Unified replies", "kind": "navigation"},
            {"id": "open:settings", "title": "Open Settings", "subtitle": "Аккаунты и настройки", "kind": "navigation"},
            {"id": "action:create_campaign", "title": "Create Campaign", "subtitle": "Wizard, no autosend", "kind": "action"},
            {"id": "action:start_session", "title": "Start Outreach Session", "subtitle": "Operator-confirmed workflow", "kind": "action"},
            {"id": "action:ai_generation", "title": "Start AI Draft Generation", "subtitle": "Drafts only, pending review", "kind": "action"},
            {"id": "filter:manual_assist", "title": "Filter Manual Assist", "subtitle": "Instagram, TikTok, X, VK", "kind": "filter"},
        ]
        if needle:
            for result in self.global_search(needle, campaign_id=campaign_id, limit=limit):
                items.append(
                    {
                        "id": f"search:{result.get('result_type')}:{result.get('object_id')}",
                        "title": str(result.get("title") or "Result"),
                        "subtitle": str(result.get("snippet") or ""),
                        "kind": str(result.get("result_type") or "search"),
                        "campaign_id": result.get("campaign_id"),
                        "contact_id": result.get("contact_id"),
                    }
                )
        else:
            for campaign in self.campaigns()[:8]:
                items.append(
                    {
                        "id": f"campaign:{campaign['id']}",
                        "title": str(campaign.get("name") or "Campaign"),
                        "subtitle": str(campaign.get("status") or "active"),
                        "kind": "campaign",
                        "campaign_id": int(campaign["id"]),
                    }
                )
            for contact in self.contacts(target_campaign)[:8]:
                items.append(
                    {
                        "id": f"contact:{contact['id']}",
                        "title": str(contact.get("company") or contact.get("name") or contact.get("email") or contact.get("handle") or "Contact"),
                        "subtitle": f"{contact.get('channel') or 'email'} • {contact.get('status') or 'new'}",
                        "kind": "contact",
                        "campaign_id": int(contact["campaign_id"]),
                        "contact_id": int(contact["id"]),
                    }
                )
        lower = needle.lower()
        if lower:
            items = [
                item for item in items
                if lower in str(item.get("title", "")).lower()
                or lower in str(item.get("subtitle", "")).lower()
                or item.get("kind") not in {"navigation", "action", "filter"}
            ]
        return items[:limit]

    def execute_command_palette_item(self, command_id: str, campaign_id: int | None = None) -> dict[str, Any]:
        command = (command_id or "").strip()
        if command == "action:start_session":
            snapshot = self.start_outreach_session(campaign_id or self.default_campaign_id())
            return {"ok": True, "message": "Outreach Session started. Autosend off.", "snapshot": snapshot, "autosend": False}
        if command == "action:ai_generation":
            return {"ok": True, "message": "AI generation opens draft flow only. Human review required.", "autosend": False}
        if command == "action:create_campaign":
            new_id = self.create_campaign_from_preset("b2b_email_outreach")
            return {"ok": True, "message": "Campaign created from preset. No sends queued.", "campaign_id": new_id, "autosend": False}
        if command.startswith("campaign:"):
            return {"ok": True, "message": "Campaign selected.", "campaign_id": int(command.split(":", 1)[1]), "autosend": False}
        return {"ok": True, "message": "Command selected. No message was sent.", "autosend": False}

    def notification_center(self, campaign_id: int | None = None, limit: int = 30) -> list[dict[str, Any]]:
        campaign_filter = "AND campaign_id = ?" if campaign_id is not None else ""
        params: list[Any] = [campaign_id] if campaign_id is not None else []
        notifications: list[dict[str, Any]] = []
        for row in self.db.fetch_all(
            f"""
            SELECT 'reply' AS type, contact_id, campaign_id, reply_status AS status,
                   substr(reply_text, 1, 180) AS message, created_at
            FROM replies
            WHERE reply_status IN ('interested', 'maybe_later', 'follow_up_needed', 'no_response') {campaign_filter}
            ORDER BY created_at DESC
            LIMIT ?
            """,
            [*params, limit],
        ):
            notifications.append(dict(row))
        for row in self.db.fetch_all(
            f"""
            SELECT 'queue_failure' AS type, contact_id, campaign_id, status,
                   COALESCE(last_error, job_type) AS message, updated_at AS created_at
            FROM job_queue
            WHERE status = 'failed' {campaign_filter}
            ORDER BY updated_at DESC
            LIMIT ?
            """,
            [*params, limit],
        ):
            notifications.append(dict(row))
        for row in self.db.fetch_all(
            f"""
            SELECT 'follow_up' AS type, contact_id, campaign_id, status,
                   note AS message, due_at AS created_at
            FROM followups
            WHERE status IN ('scheduled', 'queued') {campaign_filter}
            ORDER BY due_at ASC
            LIMIT ?
            """,
            [*params, limit],
        ):
            notifications.append(dict(row))
        for row in self.db.fetch_all(
            f"""
            SELECT 'ai_warning' AS type, id AS contact_id, campaign_id, status,
                   COALESCE(ai_warnings, research_warnings, last_error) AS message, updated_at AS created_at
            FROM contacts
            WHERE (COALESCE(ai_warnings, '') <> '' OR COALESCE(research_warnings, '') <> '' OR status = 'failed')
              {campaign_filter}
            ORDER BY updated_at DESC
            LIMIT ?
            """,
            [*params, limit],
        ):
            notifications.append(dict(row))
        notifications.sort(key=lambda row: str(row.get("created_at") or ""), reverse=True)
        return notifications[:limit]

    def background_task_monitor(self, campaign_id: int | None = None) -> dict[str, Any]:
        params: list[Any] = []
        campaign_clause = ""
        if campaign_id is not None:
            campaign_clause = "WHERE campaign_id = ?"
            params.append(campaign_id)
        rows = self.db.fetch_all(
            f"""
            SELECT status, job_type, COUNT(*) AS count
            FROM job_queue
            {campaign_clause}
            GROUP BY status, job_type
            ORDER BY status, job_type
            """,
            params,
        )
        by_status: dict[str, int] = {}
        by_type: dict[str, int] = {}
        for row in rows:
            by_status[str(row["status"])] = by_status.get(str(row["status"]), 0) + int(row["count"])
            by_type[str(row["job_type"])] = by_type.get(str(row["job_type"]), 0) + int(row["count"])
        failed = self.db.fetch_all(
            f"""
            SELECT *
            FROM job_queue
            {'WHERE campaign_id = ? AND status = ?' if campaign_id is not None else 'WHERE status = ?'}
            ORDER BY updated_at DESC, id DESC
            LIMIT 10
            """,
            ([campaign_id, "failed"] if campaign_id is not None else ["failed"]),
        )
        return {
            "queue_health": "Needs attention" if by_status.get("failed", 0) else "Healthy",
            "by_status": by_status,
            "by_type": by_type,
            "failed_tasks": failed,
            "retries_pending": by_status.get("queued", 0),
        }

    def performance_snapshot(self, campaign_id: int | None = None) -> dict[str, Any]:
        target_campaign = campaign_id or self.default_campaign_id()
        total = self.db.get_stats(target_campaign).get("total", 0)
        return {
            "campaign_id": target_campaign,
            "render_budget": recommended_batch_size(total).to_dict(),
            "cache": self.cache.stats(),
            "latency": self.latency_metrics.all_summaries(),
            "session_state": self.restore_operator_ui_state(),
        }

    def operator_global_search(self, campaign_id: int, query: str) -> list[dict[str, Any]]:
        return self.global_search(query, campaign_id=campaign_id, limit=25)

    def global_search(self, query: str, campaign_id: int | None = None, limit: int = 50) -> list[dict[str, Any]]:
        cache_key = f"search:{campaign_id or 'all'}:{limit}:{(query or '').strip().lower()}"
        cached = self.cache.get(cache_key)
        if cached is not None:
            return list(cached)  # type: ignore[arg-type]
        with self.latency_metrics.measure("global_search"):
            results = self.search_service.search(query, campaign_id=campaign_id, limit=limit)
        self.cache.set(cache_key, results, ttl_seconds=20)
        self.record_recent_search(query)
        return results

    def recent_searches(self, limit: int = 8) -> list[str]:
        try:
            values = json.loads(self.settings().get("recent_searches_json", "[]"))
        except json.JSONDecodeError:
            values = []
        return [str(value) for value in values if str(value).strip()][:limit]

    def record_recent_search(self, query: str, limit: int = 8) -> None:
        cleaned = (query or "").strip()
        if not cleaned:
            return
        existing = [item for item in self.recent_searches(limit=limit + 4) if item.lower() != cleaned.lower()]
        values = [cleaned, *existing][:limit]
        self.save_settings({"recent_searches_json": json.dumps(values, ensure_ascii=False)})

    def campaign_timeline(self, campaign_id: int, limit: int = 20) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            """
            SELECT
                ce.created_at,
                c.email,
                c.handle,
                c.external_id,
                ce.title,
                ce.event_type,
                ce.details
            FROM contact_events ce
            LEFT JOIN contacts c ON c.id = ce.contact_id
            WHERE ce.campaign_id = ?
            ORDER BY ce.created_at DESC, ce.id DESC
            LIMIT ?
            """,
            (campaign_id, limit),
        )

    def settings(self) -> dict[str, str]:
        return self.db.get_settings()

    def campaign_dashboard(self, campaign_id: int) -> dict[str, Any]:
        return self.analytics.campaign_dashboard(campaign_id)

    def campaign_metrics(self, campaign_id: int) -> dict[str, Any]:
        return self.metrics.campaign_metrics(campaign_id)

    def contact_timeline(self, contact_id: int) -> list[dict[str, Any]]:
        return self.timeline.events_for_contact(contact_id)

    def add_reply(
        self,
        contact_id: int,
        reply_text: str,
        *,
        reply_status: str = "no_response",
    ) -> int:
        suggestions = self.ai_quality.suggest_replies(contact_id, reply_text)
        return self.replies.add_reply(
            contact_id,
            reply_text,
            reply_status=reply_status,
            ai_summary=suggestions.summary,
            suggested_next_action=suggestions.suggested_next_action,
        )

    def reply_suggestions(self, contact_id: int, reply_text: str):
        return self.ai_quality.suggest_replies(contact_id, reply_text)

    def list_replies(self, campaign_id: int | None = None) -> list[dict[str, Any]]:
        return self.replies.list_replies(campaign_id)

    def unified_inbox(
        self,
        campaign_id: int | None = None,
        *,
        channel: str | None = None,
        lead_status: str | None = None,
        unread: str | None = None,
        search: str | None = None,
    ) -> list[dict[str, Any]]:
        return self.conversations.list_inbox(
            campaign_id,
            channel=channel,
            lead_status=lead_status,
            unread=unread,
            search=search,
        )

    def conversation_thread(self, contact_id: int) -> dict[str, Any]:
        return self.conversations.get_or_create_thread(contact_id)

    def conversation_messages(self, thread_id: int) -> list[dict[str, Any]]:
        return self.conversations.messages_for_thread(thread_id)

    def add_manual_reply_to_conversation(
        self,
        contact_id: int,
        reply_text: str,
        *,
        reply_status: str = "no_response",
    ) -> dict[str, Any]:
        return self.conversations.ingest_manual_reply(
            contact_id,
            reply_text,
            reply_status=reply_status,
        )

    def change_lead_status(self, contact_id: int, lead_status: str) -> str:
        return self.conversations.change_lead_status(contact_id, lead_status)

    def summarize_conversation(self, thread_id: int) -> dict[str, Any]:
        return self.conversations.summarize_thread(thread_id)

    def conversation_reply_suggestions(self, contact_id: int, reply_text: str):
        return self.conversations.suggest_replies(contact_id, reply_text)

    def suggest_followup_for_contact(self, contact_id: int) -> dict[str, Any]:
        return self.conversations.suggest_followup(contact_id)

    def conversation_metrics(self, campaign_id: int) -> dict[str, Any]:
        return self.conversations.conversation_metrics(campaign_id)

    def inbox_sync_state(self, channel: str, account_id: str = "") -> dict[str, Any]:
        return self.inbox.sync_state(channel, account_id)

    def sync_email_replies(self, campaign_id: int, client: Any | None = None, limit: int = 25) -> InboxSyncResult:
        sender = self.active_sender_email()
        password = self.active_sender_password()
        state = self.inbox.sync_state("email", sender)
        sync_client = client or EmailSyncClient()
        try:
            messages = sync_client.fetch_since(
                username=sender,
                password=password,
                last_uid=str(state.get("last_synced_uid") or ""),
                limit=limit,
            )
            result = self.inbox.ingest_email_messages(messages, campaign_id=campaign_id, account_id=sender)
        except Exception as exc:
            error = redact_secret(exc, extra_secrets=[password])
            self.inbox.save_sync_state("email", sender, status="failed", last_error=error)
            result = InboxSyncResult(ok=False, errors=[error])
        self.db.log_send(None, campaign_id, "email_sync", "ok" if result.ok else "failed", result.message, channel="email")
        return result

    def sync_telegram_replies(self, campaign_id: int, client: Any | None = None, limit: int = 50) -> InboxSyncResult:
        token = self.telegram_bot_token()
        account_id = "telegram_bot"
        state = self.inbox.sync_state("telegram", account_id)
        sync_client = client or TelegramSyncClient()
        try:
            offset = int(state.get("last_update_id") or 0) + 1
            updates = sync_client.fetch_updates(bot_token=token, offset=offset, limit=limit)
            result = self.inbox.ingest_telegram_updates(updates, campaign_id=campaign_id, account_id=account_id)
        except Exception as exc:
            error = redact_secret(exc, extra_secrets=[token])
            self.inbox.save_sync_state("telegram", account_id, status="failed", last_error=error)
            result = InboxSyncResult(ok=False, errors=[error])
        self.db.log_send(None, campaign_id, "telegram_sync", "ok" if result.ok else "failed", result.message, channel="telegram")
        return result

    def enqueue_email_sync(self, campaign_id: int, limit: int = 25) -> QueueResult:
        return self.queue.enqueue_email_sync(campaign_id, account_id=self.active_sender_email(), limit=limit)

    def enqueue_telegram_sync(self, campaign_id: int, limit: int = 50) -> QueueResult:
        return self.queue.enqueue_telegram_sync(campaign_id, limit=limit)

    def score_contact_draft(self, contact_id: int):
        return self.ai_quality.score_contact_draft(contact_id)

    def schedule_followup(
        self,
        contact_id: int,
        *,
        due_at: str | None = None,
        days_from_now: int = 2,
        note: str = "",
    ) -> int:
        return self.followups.schedule_followup(
            contact_id,
            due_at=due_at,
            days_from_now=days_from_now,
            note=note,
        )

    def enqueue_ai_score_draft(self, contact_id: int) -> QueueResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return QueueResult(False, error="Contact not found.")
        return self.queue.enqueue_ai_score_draft(
            int(contact["campaign_id"]),
            contact_id,
            channel=str(contact.get("channel") or "email"),
        )

    def enqueue_ai_generate_reply(self, contact_id: int, reply_text: str) -> QueueResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return QueueResult(False, error="Contact not found.")
        return self.queue.enqueue_ai_generate_reply(
            int(contact["campaign_id"]),
            contact_id,
            reply_text,
            channel=str(contact.get("channel") or "email"),
        )

    def enqueue_ai_summarize_reply(self, contact_id: int) -> QueueResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return QueueResult(False, error="Contact not found.")
        thread = self.conversation_thread(contact_id)
        return self.queue.enqueue_ai_summarize_reply(
            int(contact["campaign_id"]),
            contact_id,
            int(thread["id"]),
            channel=str(contact.get("channel") or "email"),
        )

    def enqueue_ai_followup_suggestion(self, contact_id: int) -> QueueResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return QueueResult(False, error="Contact not found.")
        thread = self.conversation_thread(contact_id)
        return self.queue.enqueue_ai_followup_suggestion(
            int(contact["campaign_id"]),
            contact_id,
            int(thread["id"]),
            channel=str(contact.get("channel") or "email"),
        )

    def enqueue_followup_reminder(self, contact_id: int, followup_id: int | None = None) -> QueueResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return QueueResult(False, error="Contact not found.")
        return self.queue.enqueue_followup_reminder(
            int(contact["campaign_id"]),
            contact_id,
            followup_id=followup_id,
            channel=str(contact.get("channel") or "email"),
        )

    def active_channel_id(self) -> str:
        channel_id = (self.settings().get("active_channel") or "email").strip().lower()
        return channel_id if get_channel(channel_id).channel_id == channel_id else "email"

    def active_channel(self) -> BaseChannel:
        return get_channel(self.active_channel_id())

    def available_channels(self) -> list[BaseChannel]:
        return list_channels()

    def channel_options(self) -> list[tuple[str, str]]:
        return channel_options()

    def set_active_channel(self, channel_id: str) -> str:
        channel = get_channel(channel_id)
        self.save_settings({"active_channel": channel.channel_id})
        return channel.channel_id

    def channel_capability(self, channel_id: str | None = None) -> ChannelCapability:
        return get_channel_capability(channel_id or self.active_channel_id())

    def channel_readiness_matrix(self) -> list[dict[str, Any]]:
        active_profile = self.gmail_profiles.get_active_profile()
        settings = self.settings()
        rows = list_channel_readiness(
            has_email_profile=bool(active_profile and active_profile.get("email")),
            has_telegram_token=bool(self.telegram_bot_token()),
            has_telegram_chat_id=bool(settings.get("telegram_default_test_chat_id", "").strip()),
        )
        return [row.to_dict() for row in rows]

    def connector_slots(self) -> list[dict[str, Any]]:
        active_profile = self.gmail_profiles.get_active_profile()
        settings = self.settings()
        execution_modes = {
            channel_id: self.execution_mode(channel_id)
            for channel_id in ("email", "telegram", "instagram", "tiktok", "x", "vk", "whatsapp", "viber")
        }
        slots = build_connector_slots(
            has_email_profile=bool(active_profile and active_profile.get("email")),
            has_telegram_token=bool(self.telegram_bot_token()),
            has_telegram_chat_id=bool(settings.get("telegram_default_test_chat_id", "").strip()),
            execution_modes=execution_modes,
        )
        return [slot.to_dict() for slot in slots]

    def connector_slot(self, channel_id: str) -> dict[str, Any]:
        normalized = (channel_id or "email").strip().lower()
        if normalized in {"instagram", "tiktok", "x", "vk", "whatsapp", "viber"}:
            slots = build_connector_slots(
                has_email_profile=True,
                has_telegram_token=True,
                has_telegram_chat_id=True,
                execution_modes={normalized: self.execution_mode(normalized)},
            )
            for slot in slots:
                if slot.channel_id == normalized:
                    return slot.to_dict()
        for slot in self.connector_slots():
            if slot["channel_id"] == normalized:
                return slot
        return self.connector_slots()[0]

    def channel_cockpit_snapshot(self, campaign_id: int, channel_id: str) -> dict[str, Any]:
        channel_key = (channel_id or "email").strip().lower()
        slot = self.connector_slot(channel_key)
        contacts = [row for row in self.contacts(campaign_id) if str(row.get("channel") or "email") == channel_key]
        manual_sent = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM send_logs
            WHERE campaign_id = ? AND channel = ? AND action = 'manual_assist_mark_sent'
            """,
            (campaign_id, channel_key),
        )
        copied = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM send_logs
            WHERE campaign_id = ? AND channel = ? AND action = 'operator_copy_message'
            """,
            (campaign_id, channel_key),
        )
        replies = self.db.fetch_one(
            """
            SELECT COUNT(*) AS count
            FROM replies
            LEFT JOIN contacts ON contacts.id = replies.contact_id
            WHERE replies.campaign_id = ? AND contacts.channel = ?
            """,
            (campaign_id, channel_key),
        )
        current = self.next_manual_assist_contact(campaign_id, channel=channel_key) or (contacts[0] if contacts else None)
        recommendation = self.recommend_channel_for_contact(int(current["id"])) if current else None
        action = self.execution.manual_assist_action(current).to_dict() if current else None
        return {
            "campaign_id": campaign_id,
            "channel_id": channel_key,
            "slot": slot,
            "contacts_count": len(contacts),
            "manual_sent": int(manual_sent["count"] if manual_sent else 0),
            "copied": int(copied["count"] if copied else 0),
            "replies": int(replies["count"] if replies else 0),
            "current_contact": dict(current) if current else None,
            "manual_action": action,
            "recommendation": recommendation,
            "autosend": False,
        }

    def recommend_channel_for_contact(self, contact_id: int) -> dict[str, Any]:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        return recommend_channel_for_lead(contact)

    def execution_mode_options(self, channel_id: str | None = None) -> list[tuple[str, str]]:
        capability = self.channel_capability(channel_id)
        return [(mode, EXECUTION_MODE_LABELS.get(mode, mode)) for mode in capability.allowed_execution_modes]

    def execution_mode(self, channel_id: str | None = None) -> str:
        channel_key = (channel_id or self.active_channel_id()).strip().lower()
        settings = self.settings()
        configured = settings.get(f"execution_mode_{channel_key}") or settings.get("execution_mode") or DRY_RUN
        return self.execution.policy.normalize_mode(channel_key, configured)

    def set_execution_mode(self, channel_id: str | None, mode: str) -> str:
        channel_key = (channel_id or self.active_channel_id()).strip().lower()
        normalized = self.execution.policy.normalize_mode(channel_key, mode)
        self.save_settings(
            {
                "execution_mode": normalized,
                f"execution_mode_{channel_key}": normalized,
            }
        )
        return normalized

    def execution_result_preview(
        self,
        contact: dict[str, Any],
        requested_mode: str | None = None,
        *,
        send_mode: str = "dry_run",
        confirm_live_send: bool = False,
    ) -> ChannelExecutionResult:
        return self.execution.result_for_contact(
            contact,
            requested_mode or self.execution_mode(str(contact.get("channel") or self.active_channel_id())),
            send_mode=send_mode,
            confirm_live_send=confirm_live_send,
            recipient=self.platform_recipient(contact),
        )

    def manual_assist_action_for_contact(self, contact_id: int) -> ManualAssistAction:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        return self.execution.manual_assist_action(contact)

    def next_manual_assist_contact(
        self,
        campaign_id: int,
        *,
        channel: str | None = None,
        contact_ids: list[int] | None = None,
    ) -> dict[str, Any] | None:
        channel_id = (channel or self.active_channel_id()).strip().lower()
        if contact_ids:
            placeholders = ", ".join("?" for _ in contact_ids)
            rows = self.db.fetch_all(
                f"""
                SELECT *
                FROM contacts
                WHERE campaign_id = ?
                  AND id IN ({placeholders})
                  AND channel = ?
                  AND status IN ('approved', 'pending_review', 'dry_run_sent')
                ORDER BY CASE status
                    WHEN 'approved' THEN 0
                    WHEN 'pending_review' THEN 1
                    ELSE 2
                END, id ASC
                LIMIT 1
                """,
                [campaign_id, *contact_ids, channel_id],
            )
        else:
            rows = self.db.fetch_all(
                """
                SELECT *
                FROM contacts
                WHERE campaign_id = ?
                  AND channel = ?
                  AND status = 'approved'
                ORDER BY id ASC
                LIMIT 1
                """,
                (campaign_id, channel_id),
            )
        return rows[0] if rows else None

    def prepare_manual_assist(self, contact_id: int) -> ManualAssistAction:
        action = self.manual_assist_action_for_contact(contact_id)
        contact = self.db.get_contact(contact_id)
        campaign_id = int(contact.get("campaign_id") or 0) if contact else None
        self.db.log_send(
            contact_id,
            campaign_id,
            "manual_assist_prepare",
            "manual_required",
            "Manual Assist prepared. Nothing was sent automatically.",
            channel=action.channel,
            platform_recipient=action.recipient,
        )
        self.timeline.record_event(
            contact_id,
            "note_added",
            title="Manual Assist prepared",
            details="Operator must copy/open/profile/send manually. No hidden automation.",
            metadata=action.to_dict(),
        )
        return action

    def prepare_manual_assist_campaign(
        self,
        campaign_id: int,
        *,
        channel: str | None = None,
        contact_ids: list[int] | None = None,
    ) -> QueueResult:
        channel_id = (channel or self.active_channel_id()).strip().lower()
        contact = self.next_manual_assist_contact(campaign_id, channel=channel_id, contact_ids=contact_ids)
        if not contact:
            return QueueResult(False, count=0, error="Нет подтвержденных получателей для Manual Assist.")
        self.prepare_manual_assist(int(contact["id"]))
        return QueueResult(True, count=1)

    def mark_manual_assist_sent(self, contact_ids: list[int]) -> int:
        self.backup_before_risky_operation("manual_assist_sent")
        count = 0
        for contact_id in contact_ids:
            contact = self.db.get_contact(contact_id)
            if not contact:
                continue
            action = self.execution.manual_assist_action(contact)
            sent_at = datetime.now().replace(microsecond=0)
            follow_up_days = max(as_int(self.settings().get("follow_up_delay_days"), 2), 0)
            follow_up_due_at = sent_at + timedelta(days=follow_up_days)
            self.db.update_contact(
                contact_id,
                {
                    "status": "sent",
                    "last_error": "",
                    "sent_at": sent_at.isoformat(sep=" "),
                    "follow_up_due_at": follow_up_due_at.isoformat(sep=" "),
                },
            )
            self.db.log_send(
                contact_id,
                int(contact["campaign_id"]),
                "manual_assist_mark_sent",
                "sent",
                "Operator marked message as manually sent. No platform automation was used.",
                channel=action.channel,
                platform_recipient=action.recipient,
            )
            self.timeline.record_event(
                contact_id,
                "sent",
                title="Marked manually sent",
                details=f"channel={action.channel}; recipient={action.recipient}",
                metadata={"execution_mode": MANUAL_ASSIST},
            )
            self.conversations.add_message(
                contact_id,
                direction="outbound",
                subject=action.subject,
                body=action.message,
                message_type="manual_assist",
                status="manual_sent",
                metadata=action.to_dict(),
            )
            count += 1
        return count

    def mark_manual_reply_done(self, contact_id: int, body: str = "") -> None:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        text = (body or contact.get("generated_message") or contact.get("base_message") or "Manual reply sent.").strip()
        self.conversations.add_message(
            contact_id,
            direction="outbound",
            subject=str(contact.get("subject") or ""),
            body=text,
            message_type="manual_reply",
            status="manual_sent",
            metadata={"execution_mode": MANUAL_ASSIST},
        )
        self.db.log_send(
            contact_id,
            int(contact["campaign_id"]),
            "manual_reply_marked",
            "sent",
            "Operator marked reply as manually sent. No auto-reply.",
            channel=str(contact.get("channel") or "email"),
            platform_recipient=self.platform_recipient(contact),
        )
        self.timeline.record_event(
            contact_id,
            "note_added",
            title="Manual reply marked done",
            details="Reply was handled by the operator. No automation sent it.",
        )

    def export_channel_session(
        self,
        campaign_id: int,
        channel_id: str,
        *,
        format: str = "csv",
    ) -> dict[str, Any]:
        channel_key = (channel_id or self.active_channel_id()).strip().lower()
        rows = [row for row in self.contacts(campaign_id) if str(row.get("channel") or "email") == channel_key]
        safe_format = (format or "csv").strip().lower()
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self.export_dir.mkdir(parents=True, exist_ok=True)
        fields = [
            "id",
            "channel",
            "email",
            "handle",
            "profile_url",
            "external_id",
            "name",
            "company",
            "status",
            "lead_status",
            "generated_message",
            "sent_at",
            "follow_up_due_at",
        ]
        if safe_format == "json":
            path = self.export_dir / f"{channel_key}_session_export_{timestamp}.json"
            payload = [{field: row.get(field, "") for field in fields} for row in rows]
            path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        else:
            safe_format = "csv"
            path = self.export_dir / f"{channel_key}_session_export_{timestamp}.csv"
            with path.open("w", encoding="utf-8", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=fields)
                writer.writeheader()
                for row in rows:
                    writer.writerow({field: row.get(field, "") for field in fields})
        summary = (
            f"{channel_key}: {len(rows)} contacts • "
            f"{sum(1 for row in rows if row.get('status') == 'sent')} sent/manual-sent • "
            "autosend off"
        )
        self.db.log_send(
            None,
            campaign_id,
            "channel_session_export",
            "ok",
            path.name,
            channel=channel_key,
        )
        return {
            "path": str(path),
            "format": safe_format,
            "channel": channel_key,
            "count": len(rows),
            "clipboard_summary": summary,
        }

    def mark_followup_done(self, contact_id: int) -> None:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        self.db.execute(
            """
            UPDATE followup_suggestions
            SET status = 'completed', updated_at = CURRENT_TIMESTAMP
            WHERE contact_id = ? AND status = 'suggested'
            """,
            (contact_id,),
        )
        self.db.execute(
            """
            UPDATE followups
            SET status = 'completed', updated_at = CURRENT_TIMESTAMP
            WHERE contact_id = ? AND status IN ('scheduled', 'queued')
            """,
            (contact_id,),
        )
        self.timeline.record_event(
            contact_id,
            "note_added",
            title="Follow-up marked done",
            details="Operator completed follow-up manually. No message was sent automatically.",
        )
        self.db.log_send(
            contact_id,
            int(contact["campaign_id"]),
            "manual_followup_done",
            "completed",
            "Manual follow-up marked done. No autosend.",
            channel=str(contact.get("channel") or "email"),
            platform_recipient=self.platform_recipient(contact),
        )

    def web_enrichment_settings(self) -> dict[str, Any]:
        return self.enrichment.settings_from_app_settings(self.settings())

    def platform_recipient(self, contact: dict[str, Any]) -> str:
        return get_channel(str(contact.get("channel") or "email")).platform_recipient(contact)

    def validate_channel_recipient(self, contact: dict[str, Any]) -> str:
        channel = get_channel(str(contact.get("channel") or "email"))
        ok, error = channel.validate_recipient(contact)
        if not ok:
            return self._raise_recipient_guardrail(contact, error)
        return channel.platform_recipient(contact)

    def save_settings(self, values: dict[str, str]) -> None:
        safe_values = {
            key: value
            for key, value in values.items()
            if "password" not in key.lower() and "api_key" not in key.lower()
        }
        self.backup_before_risky_operation("settings")
        self.db.set_settings(safe_values)

    def save_ai_settings(
        self,
        *,
        provider: str,
        model: str,
        api_key: str = "",
        max_drafts_per_batch: int | str = 25,
    ) -> str:
        provider_name = (provider or "off").strip().lower()
        if provider_name not in {"off", "openai"}:
            raise ValueError("Unsupported AI provider.")
        settings = {
            "ai_provider": provider_name,
            "ai_model": (model or "gpt-4.1-mini").strip(),
            "ai_max_drafts_per_batch": str(max(1, as_int(max_drafts_per_batch, 25))),
        }
        backend_name = ""
        if api_key.strip():
            backend_name = save_ai_api_key(provider_name if provider_name != "off" else "openai", api_key.strip())
        self.save_settings(settings)
        self.db.log_send(None, None, "save_ai_settings", "ok", f"provider={provider_name}; backend={backend_name or 'unchanged'}")
        return backend_name or "unchanged"

    def save_ai_brain_settings(
        self,
        brain: str,
        *,
        provider: str,
        model: str,
        api_key: str = "",
    ) -> str:
        brain_name = "research" if brain == "research" else "writer"
        provider_name = normalize_brain_provider(provider)
        settings_key_prefix = "ai_research" if brain_name == "research" else "ai_writer"
        backend_name = ""
        if api_key.strip():
            backend_name = save_ai_brain_api_key(
                brain_name,
                provider_name if provider_name != "off" else "openai",
                api_key.strip(),
            )
        self.save_settings(
            {
                f"{settings_key_prefix}_provider": provider_name,
                f"{settings_key_prefix}_model": (model or "gpt-4.1-mini").strip(),
            }
        )
        self.db.log_send(
            None,
            None,
            f"save_ai_{brain_name}_brain_settings",
            "ok",
            f"provider={provider_name}; backend={backend_name or 'unchanged'}",
        )
        return backend_name or "unchanged"

    def ai_credential_status(self, provider: str = "openai") -> CredentialStatus:
        normalized = (provider or "openai").strip().lower()
        try:
            has_password = has_ai_api_key(normalized)
        except Exception:
            has_password = False
        if has_password:
            return CredentialStatus(True, "secure AI storage", "secure_storage")
        env_key = ""
        if normalized == "openai":
            from .config import load_environment
            import os

            load_environment()
            env_key = os.getenv("OPENAI_API_KEY", "").strip()
        return CredentialStatus(
            bool(env_key),
            ".env fallback" if env_key else "secure AI storage",
            "env" if env_key else "",
        )

    def ai_brain_credential_status(self, brain: str, provider: str = "openai") -> CredentialStatus:
        brain_name = "research" if brain == "research" else "writer"
        normalized = normalize_brain_provider(provider)
        if normalized == "off":
            normalized = "openai"
        try:
            if has_ai_brain_api_key(brain_name, normalized):
                return CredentialStatus(True, f"secure {brain_name} brain storage", "secure_storage")
        except Exception:
            pass

        env_key = ""
        if normalized == "openai":
            from .config import load_environment
            import os

            load_environment()
            specific_key = "OPENAI_RESEARCH_API_KEY" if brain_name == "research" else "OPENAI_WRITER_API_KEY"
            env_key = os.getenv(specific_key, "").strip() or os.getenv("OPENAI_API_KEY", "").strip()
        return CredentialStatus(
            bool(env_key),
            ".env fallback" if env_key else f"secure {brain_name} brain storage",
            "env" if env_key else "",
        )

    def save_telegram_settings(self, *, bot_token: str = "", default_test_chat_id: str = "") -> str:
        settings = {"telegram_default_test_chat_id": default_test_chat_id.strip()}
        backend_name = ""
        if bot_token.strip():
            backend_name = save_telegram_bot_token(bot_token.strip())
        self.save_settings(settings)
        self.db.log_send(
            None,
            None,
            "save_telegram_settings",
            "ok",
            f"backend={backend_name or 'unchanged'}",
            channel="telegram",
        )
        return backend_name or "unchanged"

    def telegram_credential_status(self) -> CredentialStatus:
        status = get_telegram_credential_status()
        if status.has_password:
            return status
        env_token = get_telegram_bot_token()
        if env_token:
            return CredentialStatus(True, ".env fallback", "env")
        return status

    def telegram_bot_token(self) -> str:
        return load_telegram_bot_token() or get_telegram_bot_token()

    def check_telegram_connection(self, bot_token_override: str | None = None):
        token = (bot_token_override if bot_token_override is not None else self.telegram_bot_token()).strip()
        result = TelegramChannel().check_connection({"bot_token": token})
        self.db.log_send(
            None,
            None,
            "check_telegram_connection",
            "ok" if result.ok else "failed",
            result.message,
            channel="telegram",
        )
        if result.ok:
            self.db.set_settings({"telegram_bot_label": result.message.replace("Telegram подключен: ", "")})
        return result

    def telegram_sender_label(self) -> str:
        return self.settings().get("telegram_bot_label", "Telegram Bot").strip() or "Telegram Bot"

    def check_ai_connection(
        self,
        provider: str | None = None,
        model: str | None = None,
        api_key_override: str | None = None,
    ) -> AIConnectionCheckResult:
        settings = self.settings()
        provider_name = (provider if provider is not None else settings.get("ai_provider", "off")).strip().lower()
        model_name = (model if model is not None else settings.get("ai_model", "gpt-4.1-mini")).strip()
        service = self.ai_draft_service
        if api_key_override is not None:
            service = AIService(api_key_getter=lambda: api_key_override.strip())
        result = service.check_connection(provider_name, model_name)
        self.db.log_send(None, None, "check_ai_connection", "ok" if result.ok else "failed", result.message)
        return result

    def check_ai_brain_connection(
        self,
        brain: str,
        *,
        provider: str | None = None,
        model: str | None = None,
        api_key_override: str | None = None,
    ) -> AIConnectionCheckResult:
        brain_name = "research" if brain == "research" else "writer"
        settings = self.settings()
        provider_key = "ai_research_provider" if brain_name == "research" else "ai_writer_provider"
        model_key = "ai_research_model" if brain_name == "research" else "ai_writer_model"
        provider_name = normalize_brain_provider(provider if provider is not None else settings.get(provider_key, "off"))
        model_name = (model if model is not None else settings.get(model_key, "gpt-4.1-mini")).strip()
        if provider_name == "off":
            result = AIConnectionCheckResult(False, f"{brain_name.title()} Brain выключен.")
        else:
            if brain_name == "research":
                from .ai.research_brain import OpenAIResearchBrain
                from .config import get_openai_research_api_key

                brain_client = OpenAIResearchBrain(
                    model=model_name,
                    api_key_getter=lambda: (
                        api_key_override or load_ai_brain_api_key("research", provider_name) or get_openai_research_api_key()
                    ).strip(),
                )
            else:
                from .ai.writer_brain import OpenAIWriterBrain
                from .config import get_openai_writer_api_key

                brain_client = OpenAIWriterBrain(
                    model=model_name,
                    api_key_getter=lambda: (
                        api_key_override or load_ai_brain_api_key("writer", provider_name) or get_openai_writer_api_key()
                    ).strip(),
                )
            result = brain_client.check_connection()
        self.db.log_send(
            None,
            None,
            f"check_ai_{brain_name}_brain_connection",
            "ok" if result.ok else "failed",
            result.message,
        )
        return result

    def save_gmail_credentials(self, sender_email: str, app_password: str) -> str:
        sender = sender_email.strip().lower()
        password = app_password.strip()
        if not sender:
            raise ValueError("Gmail address is required.")
        if not password:
            raise ValueError("Gmail app password is required.")
        self.backup_before_risky_operation("gmail_credentials")
        profiles = self.gmail_profiles.list_profiles()
        active = self.gmail_profiles.get_active_profile()
        if active:
            profile = self.gmail_profiles.update_profile(
                int(active["id"]),
                name=active.get("profile_name") or "Основной",
                email=sender,
                password=password,
            )
            self.gmail_profiles.set_active_profile(int(profile["id"]))
        else:
            profile = self.gmail_profiles.create_profile("Основной", sender, password, make_active=True)
        backend_name = str(profile.get("credential_backend") or "")
        legacy_backend_name = save_gmail_app_password(sender, password)
        if not backend_name:
            backend_name = legacy_backend_name
        self.db.set_setting("sender_email", sender)
        self.db.log_send(None, None, "save_gmail_credentials", "ok", f"backend={backend_name}")
        self.logger.info("Gmail credentials saved using %s", backend_name)
        return backend_name

    def gmail_credential_status(self, sender_email: str | None = None) -> CredentialStatus:
        settings = self.settings()
        sender = (sender_email or settings.get("sender_email") or "").strip().lower()
        active = self.gmail_profiles.get_active_profile()
        if active and (
            not sender_email or sender == str(active.get("email") or "").strip().lower()
        ):
            if self.gmail_profiles.has_password(int(active["id"]), str(active["email"])):
                return CredentialStatus(True, "secure profile storage", "secure_storage")
        for profile in self.gmail_profiles.list_profiles():
            if sender and sender == str(profile.get("email") or "").strip().lower():
                if profile.get("has_password"):
                    return CredentialStatus(True, "secure profile storage", "secure_storage")
        status = get_gmail_credential_status(sender)
        if status.has_password:
            return status
        env_password = get_gmail_app_password()
        if env_password:
            return CredentialStatus(True, ".env fallback", "env")
        return status

    def list_gmail_profiles(self) -> list[dict[str, Any]]:
        return self.gmail_profiles.list_profiles()

    def create_gmail_profile(self, name: str, email: str, password: str) -> dict[str, Any]:
        self.backup_before_risky_operation("gmail_profile")
        profile = self.gmail_profiles.create_profile(name, email, password, make_active=True)
        save_gmail_app_password(str(profile.get("email") or email).strip().lower(), password)
        return profile

    def update_gmail_profile(
        self,
        profile_id: int,
        *,
        name: str | None = None,
        email: str | None = None,
        password: str | None = None,
    ) -> dict[str, Any]:
        self.backup_before_risky_operation("gmail_profile")
        old_profile = self.gmail_profiles.get_profile(profile_id)
        profile = self.gmail_profiles.update_profile(
            profile_id,
            name=name,
            email=email,
            password=password,
        )
        if password:
            save_gmail_app_password(str(profile.get("email") or email or "").strip().lower(), password)
            if old_profile and str(old_profile.get("email") or "").strip().lower() != str(profile.get("email") or "").strip().lower():
                delete_gmail_app_password(str(old_profile.get("email") or ""))
        if int(profile.get("is_active") or 0):
            self.db.set_setting("sender_email", str(profile.get("email") or ""))
        return profile

    def delete_gmail_profile(self, profile_id: int) -> None:
        self.backup_before_risky_operation("gmail_profile_delete")
        profile = self.gmail_profiles.get_profile(profile_id)
        self.gmail_profiles.delete_profile(profile_id)
        if profile:
            delete_gmail_app_password(str(profile.get("email") or ""))

    def set_active_gmail_profile(self, profile_id: int) -> dict[str, Any]:
        self.backup_before_risky_operation("gmail_profile_active")
        return self.gmail_profiles.set_active_profile(profile_id)

    def active_gmail_profile(self) -> dict[str, Any] | None:
        return self.gmail_profiles.get_active_profile()

    def active_sender_email(self) -> str:
        profile = self.active_gmail_profile()
        if profile and profile.get("email"):
            return str(profile["email"]).strip().lower()
        return (self.settings().get("sender_email") or "").strip().lower()

    def active_sender_password(self) -> str:
        profile = self.active_gmail_profile()
        if profile and profile.get("email"):
            return self.gmail_profiles.load_password(int(profile["id"]), str(profile["email"]))
        return ""

    def check_gmail_profile_connection(
        self,
        profile_id: int,
        *,
        password_override: str | None = None,
        email_override: str | None = None,
    ) -> ConnectionCheckResult:
        result = self.gmail_profiles.check_profile_connection(
            profile_id,
            password_override=password_override,
            email_override=email_override,
        )
        status = "ok" if result.ok else "failed"
        self.db.log_send(None, None, "check_gmail_profile_connection", status, result.message)
        return result

    def approved_count(self, campaign_id: int, channel: str | None = None) -> int:
        channel_id = (channel or "").strip().lower()
        if channel_id:
            row = self.db.fetch_one(
                "SELECT COUNT(*) AS count FROM contacts WHERE campaign_id = ? AND status = 'approved' AND channel = ?",
                (campaign_id, channel_id),
            )
            return int(row["count"] if row else 0)
        row = self.db.fetch_one(
            "SELECT COUNT(*) AS count FROM contacts WHERE campaign_id = ? AND status = 'approved'",
            (campaign_id,),
        )
        return int(row["count"] if row else 0)

    def approved_recipients(
        self,
        campaign_id: int,
        limit: int = 5,
        channel: str | None = None,
    ) -> list[str]:
        channel_id = (channel or "").strip().lower()
        if channel_id:
            rows = self.db.fetch_all(
                """
                SELECT *
                FROM contacts
                WHERE campaign_id = ? AND status = 'approved' AND channel = ?
                ORDER BY id ASC
                LIMIT ?
                """,
                (campaign_id, channel_id, limit),
            )
        else:
            rows = self.db.fetch_all(
                """
                SELECT *
                FROM contacts
                WHERE campaign_id = ? AND status = 'approved'
                ORDER BY id ASC
                LIMIT ?
                """,
                (campaign_id, limit),
            )
        return [self.platform_recipient(row) for row in rows]

    def validate_live_recipient(
        self,
        contact: dict[str, Any],
        *,
        safe_mode: bool,
        allowed_test_recipient: str = "",
    ) -> str:
        channel = get_channel(str(contact.get("channel") or "email"))
        if channel.channel_id == "telegram":
            ok, value = TelegramChannel().validate_live_recipient(contact)
            if not ok:
                return self._raise_recipient_guardrail(contact, value)
            chat_id = value.strip()
            allowed = (allowed_test_recipient or "").strip()
            if safe_mode and allowed and chat_id != allowed:
                return self._raise_recipient_guardrail(
                    contact,
                    "Telegram chat_id не совпадает с разрешенным тестовым получателем. "
                    f"allowed_test_recipient={allowed}; telegram_chat_id={chat_id}.",
                )
            return chat_id
        if channel.channel_id != "email" or not channel.supports_live_send:
            return self._raise_recipient_guardrail(
                contact,
                "Боевая отправка для этого канала пока отключена. "
                "Можно подготовить сообщения и сделать тестовый прогон.",
            )
        recipient = str(contact.get("email") or "").strip().lower()
        if not recipient or not is_valid_email(recipient):
            return self._raise_recipient_guardrail(
                contact,
                "Не удалось определить получателя: email в строке пустой или неверный.",
            )

        allowed = (allowed_test_recipient or "").strip().lower()
        if safe_mode and allowed and recipient != allowed:
            return self._raise_recipient_guardrail(
                contact,
                "Получатель не совпадает с разрешенным тестовым адресом. "
                f"allowed_test_recipient={allowed}; recipient_email={recipient}.",
            )
        return recipient

    def _raise_recipient_guardrail(self, contact: dict[str, Any], message: str) -> str:
        contact_id = contact.get("id")
        if contact_id:
            self.db.update_contact(int(contact_id), {"last_error": message})
        raise ValueError(message)

    def log_live_send_cancelled(self, campaign_id: int, count: int) -> None:
        self.db.log_send(
            None,
            campaign_id,
            "live_send_cancelled",
            "cancelled",
            f"Operator cancelled live send. approved_count={count}",
        )

    def enqueue_generate_messages(self, campaign_id: int, channel: str | None = None) -> QueueResult:
        self.backup_before_risky_operation("enqueue_generate")
        channel_id = (channel or self.active_channel_id()).strip().lower()
        result = self.queue.enqueue_generate_messages(campaign_id, channel=channel_id)
        self.db.log_send(
            None,
            campaign_id,
            "enqueue_generate_messages",
            "ok" if result.ok else "failed",
            f"channel={channel_id}; queued={result.count}" + (f"; {result.error}" if result.error else ""),
            channel=channel_id,
        )
        return result

    def enqueue_ai_generate_drafts(
        self,
        campaign_id: int,
        *,
        contact_ids: list[int] | None = None,
        campaign_topic: str | None = None,
        tone: str | None = None,
        confirm_ai_cost: bool = False,
        channel: str | None = None,
    ) -> QueueResult:
        settings = self.settings()
        channel_id = (channel or settings.get("active_channel") or "email").strip().lower()
        generation_mode = normalize_generation_mode(settings.get("ai_generation_mode", "simple"))
        if generation_mode == "dual_brain":
            return self.enqueue_dual_brain_generate_drafts(
                campaign_id,
                contact_ids=contact_ids,
                campaign_topic=campaign_topic,
                tone=tone,
                confirm_ai_cost=confirm_ai_cost,
                channel=channel_id,
            )
        provider_name = (settings.get("ai_provider") or "off").strip().lower()
        model_name = (settings.get("ai_model") or "gpt-4.1-mini").strip()
        topic = (campaign_topic if campaign_topic is not None else settings.get("ai_campaign_topic", "")).strip()
        tone_value = (tone if tone is not None else settings.get("ai_tone", "friendly")).strip() or "friendly"
        max_batch = max(as_int(settings.get("ai_max_drafts_per_batch"), 25), 1)

        if provider_name in {"", "off"}:
            return QueueResult(False, error="AI Assist выключен. Выберите OpenAI в Аккаунты и настройки → AI Assist.")
        if provider_name == "openai" and not (load_ai_api_key("openai") or self.ai_credential_status("openai").source == "env"):
            return QueueResult(False, error="Добавьте API key в Аккаунты и настройки → AI Assist.")
        if not topic:
            return QueueResult(False, error="Введите тему рассылки для AI Assist.")

        candidate_count = self._count_ai_draft_candidates(campaign_id, contact_ids, channel=channel_id)
        if candidate_count > max_batch and not confirm_ai_cost:
            return QueueResult(
                False,
                count=0,
                error=f"AI сгенерирует {candidate_count} писем. Это может стоить денег. Подтвердите запуск.",
            )

        self.backup_before_risky_operation("enqueue_ai_generate")
        result = self.queue.enqueue_ai_generate_drafts(
            campaign_id,
            contact_ids=contact_ids,
            campaign_topic=topic,
            tone=tone_value,
            max_count=None,
            channel=channel_id,
        )
        self.db.set_settings(
            {
                "work_mode": "ai_assist",
                "ai_campaign_topic": topic,
                "ai_tone": tone_value,
                "active_channel": channel_id,
            }
        )
        self.db.log_send(
            None,
            campaign_id,
            "enqueue_ai_generate_drafts",
            "ok" if result.ok else "failed",
            f"channel={channel_id}; queued={result.count}; provider={provider_name}; model={model_name}"
            + (f"; {result.error}" if result.error else ""),
            channel=channel_id,
        )
        return result

    def enqueue_dual_brain_generate_drafts(
        self,
        campaign_id: int,
        *,
        contact_ids: list[int] | None = None,
        campaign_topic: str | None = None,
        tone: str | None = None,
        confirm_ai_cost: bool = False,
        channel: str | None = None,
    ) -> QueueResult:
        settings = self.settings()
        channel_id = (channel or settings.get("active_channel") or "email").strip().lower()
        topic = (campaign_topic if campaign_topic is not None else settings.get("ai_campaign_topic", "")).strip()
        tone_value = (tone if tone is not None else settings.get("ai_tone", "friendly")).strip() or "friendly"
        research_provider = normalize_brain_provider(settings.get("ai_research_provider", "off"))
        writer_provider = normalize_brain_provider(settings.get("ai_writer_provider", "off"))
        research_model = (settings.get("ai_research_model") or "gpt-4.1-mini").strip()
        writer_model = (settings.get("ai_writer_model") or "gpt-4.1-mini").strip()
        max_batch = max(as_int(settings.get("ai_max_drafts_per_batch"), 25), 1)

        if not topic:
            return QueueResult(False, error="Введите тему рассылки для AI Assist.")
        if research_provider in {"", "off"}:
            return QueueResult(False, error="Research Brain выключен. Выберите OpenAI в Аккаунты и настройки → AI Assist.")
        if writer_provider in {"", "off"}:
            return QueueResult(False, error="Writer Brain выключен. Выберите OpenAI в Аккаунты и настройки → AI Assist.")
        if research_provider == "openai" and not self.ai_brain_credential_status("research", "openai").has_password:
            return QueueResult(False, error="Добавьте Research Brain API key в Аккаунты и настройки → AI Assist.")
        if writer_provider == "openai" and not self.ai_brain_credential_status("writer", "openai").has_password:
            return QueueResult(False, error="Добавьте Writer Brain API key в Аккаунты и настройки → AI Assist.")

        candidate_count = self._count_ai_draft_candidates(campaign_id, contact_ids, channel=channel_id)
        if candidate_count > max_batch and not confirm_ai_cost:
            return QueueResult(
                False,
                count=0,
                error=f"AI сгенерирует {candidate_count} писем. Это может стоить денег. Подтвердите запуск.",
            )

        enrichment_enabled = self.web_enrichment_settings().get("enabled", False)
        self.backup_before_risky_operation("enqueue_ai_dual_brain_generate")
        result = self.queue.enqueue_ai_dual_brain_drafts(
            campaign_id,
            contact_ids=contact_ids,
            campaign_topic=topic,
            tone=tone_value,
            max_count=None,
            channel=channel_id,
            use_enrichment=bool(enrichment_enabled),
        )
        self.db.set_settings(
            {
                "work_mode": "ai_assist",
                "ai_generation_mode": "dual_brain",
                "ai_campaign_topic": topic,
                "ai_tone": tone_value,
                "active_channel": channel_id,
                "web_enrichment_enabled": "true" if enrichment_enabled else "false",
            }
        )
        self.db.log_send(
            None,
            campaign_id,
            "enqueue_ai_dual_brain_generate",
            "ok" if result.ok else "failed",
            (
                f"channel={channel_id}; queued={result.count}; "
                f"research={research_provider}/{research_model}; writer={writer_provider}/{writer_model}; "
                f"web_enrichment={bool(enrichment_enabled)}"
            )
            + (f"; {result.error}" if result.error else ""),
            channel=channel_id,
        )
        return result

    def _count_ai_draft_candidates(
        self,
        campaign_id: int,
        contact_ids: list[int] | None = None,
        channel: str | None = None,
    ) -> int:
        channel_clause = ""
        channel_params: list[Any] = []
        if channel:
            channel_clause = " AND channel = ?"
            channel_params.append(channel)
        if contact_ids:
            placeholders = ", ".join("?" for _ in contact_ids)
            row = self.db.fetch_one(
                f"""
                SELECT COUNT(*) AS count
                FROM contacts
                WHERE campaign_id = ?
                  AND id IN ({placeholders})
                  AND status IN ('new', 'failed')
                  {channel_clause}
                """,
                [campaign_id, *contact_ids, *channel_params],
            )
        else:
            row = self.db.fetch_one(
                f"""
                SELECT COUNT(*) AS count
                FROM contacts
                WHERE campaign_id = ?
                  AND status IN ('new', 'failed')
                  {channel_clause}
                """,
                [campaign_id, *channel_params],
            )
        return int(row["count"] if row else 0)

    def ai_draft_candidate_count(
        self,
        campaign_id: int,
        contact_ids: list[int] | None = None,
        channel: str | None = None,
    ) -> int:
        return self._count_ai_draft_candidates(campaign_id, contact_ids, channel=channel or self.active_channel_id())

    def enqueue_enrich_contacts(
        self,
        campaign_id: int,
        *,
        contact_ids: list[int] | None = None,
        channel: str | None = None,
        force_refresh: bool = False,
    ) -> QueueResult:
        settings = self.settings()
        max_batch = max(as_int(settings.get("web_enrichment_max_contacts_per_batch"), 25), 1)
        channel_id = (channel or settings.get("active_channel") or "email").strip().lower()
        self.backup_before_risky_operation("enqueue_enrich_contacts")
        result = self.queue.enqueue_enrich_contacts(
            campaign_id,
            contact_ids=contact_ids,
            max_count=max_batch,
            channel=channel_id,
            force_refresh=force_refresh,
        )
        self.db.log_send(
            None,
            campaign_id,
            "enqueue_enrich_contacts",
            "ok" if result.ok else "failed",
            f"channel={channel_id}; queued={result.count}; max_batch={max_batch}"
            + (f"; {result.error}" if result.error else ""),
            channel=channel_id,
        )
        return result

    def enrich_contact_for_research(
        self,
        contact_id: int,
        *,
        force_refresh: bool = False,
        enabled_override: bool | None = None,
    ) -> EnrichmentResult:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found.")
        settings = self.web_enrichment_settings()
        if enabled_override is not None:
            settings["enabled"] = bool(enabled_override)
        result = self.enrichment.enrich_contact(contact, settings=settings, force_refresh=force_refresh)
        self.enrichment.save_contact_result(contact_id, result)
        self.timeline.record_event(
            contact_id,
            "status_changed",
            title="Web enrichment checked",
            details=f"status={result.status}; confidence={result.confidence}",
            metadata={"source_urls": result.source_urls, "warnings": result.warnings},
        )
        self.db.log_send(
            contact_id,
            int(contact.get("campaign_id") or 0) or None,
            "enrich_contact",
            result.status,
            "; ".join(result.warnings),
            channel=str(contact.get("channel") or "email"),
            platform_recipient=self.platform_recipient(contact),
        )
        return result

    def generate_ai_draft_for_contact(
        self,
        contact: dict[str, Any],
        *,
        campaign_topic: str,
        tone: str,
    ):
        settings = self.settings()
        return self.ai_draft_service.generate_draft_for_contact(
            contact,
            campaign_topic,
            tone,
            provider_name=settings.get("ai_provider", "off"),
            model=settings.get("ai_model", "gpt-4.1-mini"),
            execution_mode=self.execution_mode(str(contact.get("channel") or self.active_channel_id())),
            execution_notes=self.channel_capability(str(contact.get("channel") or "email")).limitation_text,
        )

    def _research_brain_config(self) -> BrainConfig:
        settings = self.settings()
        return BrainConfig(
            provider=normalize_brain_provider(settings.get("ai_research_provider", "off")),
            model=(settings.get("ai_research_model") or "gpt-4.1-mini").strip(),
        )

    def _writer_brain_config(self) -> BrainConfig:
        settings = self.settings()
        return BrainConfig(
            provider=normalize_brain_provider(settings.get("ai_writer_provider", "off")),
            model=(settings.get("ai_writer_model") or "gpt-4.1-mini").strip(),
        )

    def research_contact_for_ai(
        self,
        contact: dict[str, Any],
        *,
        campaign_topic: str,
    ) -> RecipientBrief:
        return self.dual_brain_service.research_contact(
            contact,
            campaign_topic,
            config=self._research_brain_config(),
        )

    def write_draft_from_existing_brief(
        self,
        contact: dict[str, Any],
        *,
        campaign_topic: str,
        tone: str,
    ):
        raw_brief = contact.get("research_brief_json") or ""
        if not raw_brief:
            raise ValueError("Research brief отсутствует. Сначала запустите Research Brain.")
        brief = parse_recipient_brief(raw_brief)
        return self.dual_brain_service.write_from_brief(
            contact,
            campaign_topic,
            tone,
            brief,
            config=self._writer_brain_config(),
            execution_mode=self.execution_mode(str(contact.get("channel") or self.active_channel_id())),
            execution_notes=self.channel_capability(str(contact.get("channel") or "email")).limitation_text,
        )

    def generate_dual_brain_draft_for_contact(
        self,
        contact: dict[str, Any],
        *,
        campaign_topic: str,
        tone: str,
    ) -> DualBrainResult:
        return self.dual_brain_service.generate(
            contact,
            campaign_topic,
            tone,
            research_config=self._research_brain_config(),
            writer_config=self._writer_brain_config(),
            execution_mode=self.execution_mode(str(contact.get("channel") or self.active_channel_id())),
            execution_notes=self.channel_capability(str(contact.get("channel") or "email")).limitation_text,
        )

    def record_dual_brain_metrics(
        self,
        contact_id: int,
        *,
        result: DualBrainResult,
    ) -> None:
        contact = self.db.get_contact(contact_id)
        if not contact:
            return
        research_config = self._research_brain_config()
        writer_config = self._writer_brain_config()
        warnings = [*result.brief.warnings, *result.draft.warnings]
        self.db.execute(
            """
            INSERT INTO ai_metrics (
                contact_id, campaign_id, spam_risk, personalization_quality,
                confidence, genericness_score, tone_quality, warnings,
                research_model, writer_model, research_confidence, writer_confidence
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                contact_id,
                int(contact["campaign_id"]),
                "low",
                result.brief.personalization_strength,
                result.draft.confidence,
                1.0 - result.brief.confidence,
                "ok" if not result.draft.warnings else "warning",
                redact_secret("; ".join(warnings)),
                research_config.model,
                writer_config.model,
                result.brief.confidence,
                result.draft.confidence,
            ),
        )

    def enqueue_send_approved(
        self,
        campaign_id: int,
        mode: str,
        confirm_live_send: bool = False,
        channel: str | None = None,
    ) -> QueueResult:
        self.backup_before_risky_operation("enqueue_send")
        channel_id = (channel or self.active_channel_id()).strip().lower()
        channel_obj = get_channel(channel_id)
        execution_mode = self.execution_mode(channel_obj.channel_id)
        if execution_mode == MANUAL_ASSIST:
            result = self.prepare_manual_assist_campaign(campaign_id, channel=channel_obj.channel_id)
            self.db.log_send(
                None,
                campaign_id,
                "execution_manual_assist",
                "manual_required" if result.ok else "failed",
                (
                    f"channel={channel_obj.channel_id}; mode=manual_assist; prepared={result.count}; "
                    "no hidden automation"
                )
                + (f"; {result.error}" if result.error else ""),
                channel=channel_obj.channel_id,
            )
            return result
        if execution_mode == OFFICIAL_API and mode not in {"dry_run", "live"}:
            mode = "dry_run"
        if mode == "live" and channel_obj.channel_id == "telegram" and not self.telegram_bot_token():
            return QueueResult(
                False,
                count=0,
                error="Telegram live send blocked: сохраните Bot Token в Аккаунты и настройки → Telegram.",
            )
        if mode == "live" and not channel_obj.supports_live_send:
            return QueueResult(
                False,
                count=0,
                error=channel_obj.live_disabled_message,
            )
        result = self.queue.enqueue_bulk_send(
            campaign_id,
            mode=mode,
            confirm_live_send=confirm_live_send,
            channel=channel_id,
        )
        self.db.log_send(
            None,
            campaign_id,
            "enqueue_send",
            "ok" if result.ok else "failed",
            f"channel={channel_id}; mode={mode}; queued={result.count}" + (f"; {result.error}" if result.error else ""),
            channel=channel_id,
        )
        return result

    def enqueue_export_report(self, campaign_id: int) -> QueueResult:
        self.backup_before_risky_operation("enqueue_export")
        result = self.queue.enqueue_export_report(campaign_id)
        self.db.log_send(
            None,
            campaign_id,
            "enqueue_export_report",
            "ok" if result.ok else "failed",
            f"queued={result.count}" + (f"; {result.error}" if result.error else ""),
        )
        return result

    def queue_stats(self, campaign_id: int | None = None) -> dict[str, int]:
        return self.queue.get_queue_stats(campaign_id)

    def recent_jobs(self, campaign_id: int | None = None, limit: int = 50) -> list[dict[str, Any]]:
        return self.queue.list_recent_jobs(campaign_id, limit=limit)

    def retry_failed_jobs(self, campaign_id: int | None = None) -> QueueResult:
        return self.queue.retry_failed_jobs(campaign_id)

    def clear_completed_jobs(self, campaign_id: int | None = None) -> int:
        return self.queue.clear_completed_jobs(campaign_id)

    def check_gmail_connection(
        self,
        sender_email: str | None = None,
        app_password: str | None = None,
    ) -> ConnectionCheckResult:
        settings = self.settings()
        profile = self.active_gmail_profile()
        sender = (
            sender_email
            if sender_email is not None
            else (str(profile.get("email") or "") if profile else settings.get("sender_email", ""))
        ).strip()
        profile_password = ""
        matched_profile: dict[str, Any] | None = None
        for candidate in self.gmail_profiles.list_profiles():
            if sender and sender.lower() == str(candidate.get("email") or "").strip().lower():
                matched_profile = candidate
                break
        if app_password is None and matched_profile:
            profile_password = self.gmail_profiles.load_password(
                int(matched_profile["id"]),
                str(matched_profile["email"]),
            )
        result = self.mailer.check_connection(
            host=settings.get("smtp_host", "smtp.gmail.com"),
            port=as_int(settings.get("smtp_port"), 587),
            sender_email=sender,
            password=app_password if app_password is not None else (profile_password or None),
        )
        status = "ok" if result.ok else "failed"
        if matched_profile:
            self.gmail_profiles._update_check_status(int(matched_profile["id"]), result)
        self.db.log_send(None, None, "check_gmail_connection", status, result.message)
        if result.ok:
            self.logger.info("Gmail connection check succeeded")
        else:
            self.logger.warning("Gmail connection check failed: %s", redact_secret(result.message))
        return result

    def template(self) -> dict[str, Any]:
        return self.db.get_template("Default")

    def save_template(self, subject_template: str, body_template: str) -> None:
        self.backup_before_risky_operation("template")
        self.db.save_template("Default", subject_template, body_template)

    def import_file(self, file_path: str | Path, campaign_id: int) -> ImportResult:
        self.backup_before_risky_operation("import")
        result = self.importer.import_file(file_path, campaign_id)
        self.db.log_send(
            None,
            campaign_id,
            "import_contacts",
            "ok",
            f"imported={result.imported_count}; skipped={result.skipped_count}",
        )
        return result

    def add_contact_row(self, campaign_id: int, row: dict[str, Any]) -> ImportResult:
        return self.add_contact_rows(campaign_id, [row], source="manual")

    def add_contact_rows(
        self,
        campaign_id: int,
        rows: list[dict[str, Any]],
        source: str = "manual",
        *,
        create_backup: bool = True,
    ) -> ImportResult:
        if create_backup:
            self.backup_before_risky_operation("contacts")
        result = ImportResult()
        seen: set[str] = set()
        active_channel = self.active_channel_id()
        for index, row in enumerate(rows, start=1):
            channel_id = str(row.get("channel") or active_channel or "email").strip().lower()
            channel = get_channel(channel_id)
            channel_id = channel.channel_id
            handle = str(row.get("handle") or "").strip()
            profile_url = str(row.get("profile_url") or "").strip()
            external_id = str(row.get("external_id") or "").strip()
            email = str(row.get("email") or "").strip().lower()

            if channel_id == "email" and (not email or not is_valid_email(email)):
                result.skipped_count += 1
                message = "Неверный или пустой email"
                result.errors.append(f"Строка {index}: {message}")
                self.db.log_import_error(campaign_id, source, index, email, message)
                continue

            platform_recipient = channel.platform_recipient(
                {
                    "email": email,
                    "handle": handle,
                    "profile_url": profile_url,
                    "external_id": external_id,
                }
            )
            if channel_id != "email" and not platform_recipient:
                result.skipped_count += 1
                message = f"Не указан получатель для канала {channel.display_name}"
                result.errors.append(f"Строка {index}: {message}")
                self.db.log_import_error(campaign_id, source, index, email, message)
                continue

            storage_email = email
            if channel_id != "email" and (not storage_email or not is_valid_email(storage_email)):
                storage_email = self._synthetic_channel_email(channel_id, platform_recipient)

            duplicate_key = f"{channel_id}:{platform_recipient or storage_email}".lower()
            if duplicate_key in seen or self._contact_exists_for_channel(
                campaign_id,
                channel_id,
                platform_recipient,
                storage_email,
            ):
                result.skipped_count += 1
                message = "Дубликат получателя в этой рассылке"
                result.errors.append(f"Строка {index}: {message} ({platform_recipient or storage_email})")
                self.db.log_import_error(campaign_id, source, index, storage_email, message)
                continue
            seen.add(duplicate_key)
            message_text = str(
                row.get("generated_message") or row.get("base_message") or row.get("message") or ""
            ).strip()
            last_error = (
                ""
                if message_text
                else "Нет сообщения. Можно заполнить вручную или использовать шаблон."
            )
            contact_id = self.db.add_contact(
                {
                    "campaign_id": campaign_id,
                    "channel": channel_id,
                    "email": storage_email,
                    "handle": handle,
                    "profile_url": profile_url,
                    "external_id": external_id,
                    "name": str(row.get("name") or "").strip(),
                    "company": str(row.get("company") or "").strip(),
                    "topic": str(row.get("topic") or "").strip(),
                    "website": str(row.get("website") or "").strip(),
                    "social_profile": str(row.get("social_profile") or "").strip(),
                    "subject": str(row.get("subject") or "").strip(),
                    "base_message": message_text,
                    "generated_message": message_text,
                    "last_error": last_error,
                    "status": str(row.get("status") or "new").strip() or "new",
                }
            )
            self.timeline.record_event(
                contact_id,
                "contact_created",
                title="Контакт создан",
                details=f"channel={channel_id}; recipient={platform_recipient or storage_email}",
                metadata={"channel": channel_id},
            )
            self.conversations.get_or_create_thread(contact_id)
            result.imported_count += 1
        self.db.log_send(
            None,
            campaign_id,
            "add_contacts",
            "ok" if result.imported_count else "failed",
            f"imported={result.imported_count}; skipped={result.skipped_count}",
        )
        return result

    def _synthetic_channel_email(self, channel_id: str, platform_recipient: str) -> str:
        digest = hashlib.sha1(f"{channel_id}:{platform_recipient}".encode("utf-8")).hexdigest()[:14]
        return f"{channel_id}-{digest}@channel.local"

    def _contact_exists_for_channel(
        self,
        campaign_id: int,
        channel_id: str,
        platform_recipient: str,
        storage_email: str,
    ) -> bool:
        if channel_id == "email":
            return self.db.get_contact_by_email(campaign_id, storage_email) is not None
        row = self.db.fetch_one(
            """
            SELECT id
            FROM contacts
            WHERE campaign_id = ?
              AND channel = ?
              AND (
                lower(email) = lower(?)
                OR lower(handle) = lower(?)
                OR lower(profile_url) = lower(?)
                OR lower(external_id) = lower(?)
              )
            LIMIT 1
            """,
            (
                campaign_id,
                channel_id,
                storage_email,
                platform_recipient,
                platform_recipient,
                platform_recipient,
            ),
        )
        return row is not None

    def delete_contacts(self, contact_ids: list[int]) -> int:
        self.backup_before_risky_operation("delete_contacts")
        return self.db.delete_contacts(contact_ids)

    def export_example_contacts_template(self) -> Path:
        self.export_dir.mkdir(parents=True, exist_ok=True)
        path = self.export_dir / "example_contacts_template.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Получатели"
        sheet.append(
            [
                "Email",
                "Тема письма",
                "Сообщение",
                "Имя",
                "Компания",
                "Сайт",
                "Соцсеть / профиль",
                "Заметка",
                "Канал",
                "Профиль / username",
                "URL профиля",
                "ID / chat",
            ]
        )
        sheet.append(
            [
                "first@example.com",
                "Сотрудничество",
                "Здравствуйте! Хочу предложить аккуратный тестовый сценарий.",
                "Анна",
                "Example Co",
                "https://example.com",
                "https://linkedin.com/company/example",
                "Безопасный пример",
                "email",
                "",
                "",
                "",
            ]
        )
        sheet.append(
            [
                "",
                "",
                "Здравствуйте! Это безопасный пример social dry-run без live-отправки.",
                "Иван",
                "Example Media",
                "https://example.org",
                "https://x.com/example_profile",
                "Канал X пока только dry-run",
                "x",
                "@example_profile",
                "https://x.com/example_profile",
                "",
            ]
        )
        sheet.append(
            [
                "second@example.com",
                "Вопрос по рекламе",
                "Добрый день! Интересует возможность обсудить размещение.",
                "Иван",
                "Example Media",
                "https://example.org",
                "",
                "Еще один example.com адрес",
                "email",
                "",
                "",
                "",
            ]
        )
        workbook.save(path)
        self.db.log_send(None, None, "export_example_template", "ok", path.name)
        return path

    def massive_import_rows(
        self,
        campaign_id: int,
        rows: list[dict[str, Any]],
        *,
        chunk_size: int = 500,
        source: str = "massive_import",
    ) -> dict[str, Any]:
        self.backup_before_risky_operation("massive_import")
        summary = ImportResult()
        chunks = len(list(chunked(rows, chunk_size)))
        active_channel = self.active_channel_id()
        existing_rows = self.db.fetch_all(
            "SELECT channel, email, handle, profile_url, external_id FROM contacts WHERE campaign_id = ?",
            (campaign_id,),
        )
        existing_keys: set[str] = set()
        for existing in existing_rows:
            channel_id = str(existing.get("channel") or "email").strip().lower()
            for field in ("email", "handle", "profile_url", "external_id"):
                value = str(existing.get(field) or "").strip().lower()
                if value:
                    existing_keys.add(f"{channel_id}:{value}")
        seen: set[str] = set()

        with self.db.connect() as conn:
            for index, row in enumerate(rows, start=1):
                channel_id = str(row.get("channel") or active_channel or "email").strip().lower()
                channel = get_channel(channel_id)
                channel_id = channel.channel_id
                handle = str(row.get("handle") or "").strip()
                profile_url = str(row.get("profile_url") or "").strip()
                external_id = str(row.get("external_id") or "").strip()
                email = str(row.get("email") or "").strip().lower()

                if channel_id == "email" and (not email or not is_valid_email(email)):
                    summary.skipped_count += 1
                    message = "Неверный или пустой email"
                    summary.errors.append(f"Строка {index}: {message}")
                    conn.execute(
                        "INSERT INTO import_errors (campaign_id, source_file, row_number, email, error) VALUES (?, ?, ?, ?, ?)",
                        (campaign_id, source, index, email, redact_secret(message)),
                    )
                    continue

                platform_recipient = channel.platform_recipient(
                    {
                        "email": email,
                        "handle": handle,
                        "profile_url": profile_url,
                        "external_id": external_id,
                    }
                )
                if channel_id != "email" and not platform_recipient:
                    summary.skipped_count += 1
                    message = f"Не указан получатель для канала {channel.display_name}"
                    summary.errors.append(f"Строка {index}: {message}")
                    conn.execute(
                        "INSERT INTO import_errors (campaign_id, source_file, row_number, email, error) VALUES (?, ?, ?, ?, ?)",
                        (campaign_id, source, index, email, redact_secret(message)),
                    )
                    continue

                storage_email = email
                if channel_id != "email" and (not storage_email or not is_valid_email(storage_email)):
                    storage_email = self._synthetic_channel_email(channel_id, platform_recipient)
                duplicate_key = f"{channel_id}:{platform_recipient or storage_email}".lower()
                lookup_values = {
                    f"{channel_id}:{storage_email.lower()}",
                    f"{channel_id}:{handle.lower()}",
                    f"{channel_id}:{profile_url.lower()}",
                    f"{channel_id}:{external_id.lower()}",
                    duplicate_key,
                }
                if duplicate_key in seen or any(value in existing_keys for value in lookup_values if value != f"{channel_id}:"):
                    summary.skipped_count += 1
                    message = "Дубликат получателя в этой рассылке"
                    summary.errors.append(f"Строка {index}: {message} ({platform_recipient or storage_email})")
                    conn.execute(
                        "INSERT INTO import_errors (campaign_id, source_file, row_number, email, error) VALUES (?, ?, ?, ?, ?)",
                        (campaign_id, source, index, storage_email, redact_secret(message)),
                    )
                    continue
                seen.add(duplicate_key)
                existing_keys.update(value for value in lookup_values if value != f"{channel_id}:")

                message_text = str(
                    row.get("generated_message") or row.get("base_message") or row.get("message") or ""
                ).strip()
                last_error = "" if message_text else "Нет сообщения. Можно заполнить вручную или использовать шаблон."
                cursor = conn.execute(
                    """
                    INSERT INTO contacts (
                        campaign_id, channel, email, handle, profile_url, external_id,
                        name, company, topic, website, social_profile, subject,
                        base_message, generated_message, status, last_error
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        campaign_id,
                        channel_id,
                        storage_email,
                        handle,
                        profile_url,
                        external_id,
                        str(row.get("name") or "").strip(),
                        str(row.get("company") or "").strip(),
                        str(row.get("topic") or "").strip(),
                        str(row.get("website") or "").strip(),
                        str(row.get("social_profile") or "").strip(),
                        str(row.get("subject") or "").strip(),
                        message_text,
                        message_text,
                        str(row.get("status") or "new").strip() or "new",
                        last_error,
                    ),
                )
                contact_id = int(cursor.lastrowid)
                conn.execute(
                    """
                    INSERT INTO contact_events (
                        contact_id, campaign_id, event_type, title, details, metadata_json
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        contact_id,
                        campaign_id,
                        "contact_created",
                        "Контакт создан",
                        f"channel={channel_id}; recipient={platform_recipient or storage_email}",
                        json.dumps({"channel": channel_id}, ensure_ascii=False),
                    ),
                )
                thread_cursor = conn.execute(
                    """
                    INSERT INTO conversation_threads (
                        campaign_id, contact_id, channel, lead_status, unread_state
                    ) VALUES (?, ?, ?, ?, ?)
                    """,
                    (
                        campaign_id,
                        contact_id,
                        channel_id,
                        str(row.get("lead_status") or "New"),
                        "read",
                    ),
                )
                conn.execute(
                    """
                    INSERT INTO contact_events (
                        contact_id, campaign_id, event_type, title, details, metadata_json
                    ) VALUES (?, ?, ?, ?, ?, ?)
                    """,
                    (
                        contact_id,
                        campaign_id,
                        "note_added",
                        "Conversation thread created",
                        "Conversation timeline initialized.",
                        json.dumps({"thread_id": int(thread_cursor.lastrowid)}, ensure_ascii=False),
                    ),
                )
                summary.imported_count += 1
        self.db.log_send(
            None,
            campaign_id,
            "massive_import",
            "ok" if summary.imported_count else "failed",
            f"chunks={chunks}; imported={summary.imported_count}; skipped={summary.skipped_count}; no autosend",
        )
        self.cache.invalidate("search:")
        return {
            "imported_count": summary.imported_count,
            "skipped_count": summary.skipped_count,
            "errors": summary.errors,
            "chunks": chunks,
            "autosend": False,
        }

    def contacts(
        self,
        campaign_id: int,
        status: str | None = None,
        search: str | None = None,
    ) -> list[dict[str, Any]]:
        return self.db.list_contacts(campaign_id, status=status, search=search)

    def contacts_page(
        self,
        campaign_id: int,
        *,
        offset: int = 0,
        limit: int = 250,
        status: str | None = None,
        search: str | None = None,
    ) -> dict[str, Any]:
        clauses = ["campaign_id = ?"]
        params: list[Any] = [campaign_id]
        if status and status != "all":
            clauses.append("status = ?")
            params.append(status)
        if search:
            pattern = f"%{search.strip()}%"
            clauses.append(
                "(email LIKE ? OR name LIKE ? OR company LIKE ? OR topic LIKE ? "
                "OR website LIKE ? OR social_profile LIKE ? OR channel LIKE ? "
                "OR handle LIKE ? OR profile_url LIKE ? OR external_id LIKE ?)"
            )
            params.extend([pattern] * 10)
        where = " AND ".join(clauses)
        count_row = self.db.fetch_one(f"SELECT COUNT(*) AS count FROM contacts WHERE {where}", params)
        total = int(count_row["count"] if count_row else 0)
        page = virtual_page(total, offset=offset, limit=limit)
        items = self.db.fetch_all(
            f"""
            SELECT *
            FROM contacts
            WHERE {where}
            ORDER BY id DESC
            LIMIT ? OFFSET ?
            """,
            [*params, page.limit, page.offset],
        )
        return {
            "items": items,
            "page": page.to_dict(),
            "render_budget": recommended_batch_size(total).to_dict(),
        }

    def export_operator_data(self, campaign_id: int, kind: str = "leads", *, format: str = "csv") -> dict[str, Any]:
        export_kind = (kind or "leads").strip().lower()
        export_format = (format or "csv").strip().lower()
        if export_format not in {"csv", "json"}:
            raise ValueError("Unsupported export format.")
        self.export_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = self.export_dir / f"operator_{export_kind}_{campaign_id}_{timestamp}.{export_format}"

        if export_kind == "inbox":
            rows = self.unified_inbox(campaign_id=campaign_id)[:1000]
        elif export_kind == "analytics":
            rows = [self.operator_analytics_summary(campaign_id), self.performance_snapshot(campaign_id)]
        else:
            rows = self.contacts(campaign_id)

        if export_format == "json":
            path.write_text(json.dumps(rows, ensure_ascii=False, indent=2, default=str), encoding="utf-8")
        else:
            with path.open("w", newline="", encoding="utf-8") as handle:
                if not rows:
                    handle.write("")
                else:
                    keys = sorted({key for row in rows if isinstance(row, dict) for key in row.keys()})
                    writer = csv.DictWriter(handle, fieldnames=keys)
                    writer.writeheader()
                    for row in rows:
                        if isinstance(row, dict):
                            writer.writerow({key: row.get(key, "") for key in keys})
        self.db.log_send(None, campaign_id, f"export_{export_kind}", "ok", path.name)
        return {"path": str(path), "kind": export_kind, "format": export_format, "rows": len(rows), "autosend": False}

    def stats(self, campaign_id: int) -> dict[str, int]:
        return self.db.get_stats(campaign_id)

    def update_contact(self, contact_id: int, fields: dict[str, Any]) -> None:
        self.db.update_contact(contact_id, fields)

    def generate_messages(self, campaign_id: int) -> int:
        settings = self.settings()
        review_mode = as_bool(settings.get("review_mode"), True)
        template = self.template()
        contacts = self.db.list_contacts(campaign_id, status="new")
        next_status = "pending_review" if review_mode else "approved"
        generated_count = 0

        for contact in contacts:
            generated = self.ai_provider.generate_email(
                contact,
                template["subject_template"],
                template["body_template"],
            )
            self.db.update_contact(
                contact["id"],
                {
                    "subject": generated.subject,
                    "generated_message": generated.body,
                    "status": next_status,
                    "last_error": "",
                },
            )
            generated_count += 1

        self.db.log_send(
            None,
            campaign_id,
            "generate_messages",
            "ok",
            f"generated={generated_count}; status={next_status}",
        )
        return generated_count

    def preview_generated_email(self, contact_id: int) -> tuple[str, str]:
        contact = self.db.get_contact(contact_id)
        if not contact:
            raise ValueError("Contact not found")
        template = self.template()
        generated = self.ai_provider.generate_email(
            contact,
            template["subject_template"],
            template["body_template"],
        )
        return generated.subject, generated.body

    def approve_contacts(self, contact_ids: list[int]) -> int:
        self.backup_before_risky_operation("approve_contacts")
        count = 0
        for contact_id in contact_ids:
            before = self.db.get_contact(contact_id)
            self.db.update_contact(contact_id, {"status": "approved", "last_error": ""})
            self.timeline.record_event(
                contact_id,
                "approved",
                title="Письмо подтверждено",
                details="Human review completed.",
            )
            if before and before.get("status") != "approved":
                self.timeline.record_status_change(contact_id, str(before.get("status") or ""), "approved")
            count += 1
        return count

    def blacklist_contacts(self, contact_ids: list[int], reason: str = "Manual blacklist") -> int:
        self.backup_before_risky_operation("blacklist")
        count = 0
        for contact_id in contact_ids:
            self.blacklist.blacklist_contact(contact_id, reason)
            contact = self.db.get_contact(contact_id)
            self.db.log_send(
                contact_id,
                contact["campaign_id"] if contact else None,
                "blacklist",
                "blacklisted",
                reason,
            )
            self.timeline.record_event(
                contact_id,
                "blacklisted",
                title="Добавлен в blacklist",
                details=reason,
            )
            count += 1
        return count

    def send_approved(self, campaign_id: int, confirm_live_send: bool = False) -> SendSummary:
        settings = self.settings()
        send_mode = (settings.get("send_mode") or "dry_run").strip().lower()
        if send_mode not in {"dry_run", "live"}:
            send_mode = "dry_run"
        host = settings.get("smtp_host", "smtp.gmail.com")
        port = as_int(settings.get("smtp_port"), 587)
        sender_email = self.active_sender_email()
        sender_password = self.active_sender_password()
        daily_limit = as_int(settings.get("daily_send_limit"), 25)
        delay_seconds = max(as_int(settings.get("delay_seconds"), 0), 0)
        follow_up_days = max(as_int(settings.get("follow_up_delay_days"), 2), 0)
        safe_mode = as_bool(settings.get("safe_mode"), True)
        real_send_confirm_required = as_bool(
            settings.get("real_send_confirm_required"),
            True,
        )
        allowed_test_recipient = (settings.get("allowed_test_recipient") or "").strip()

        summary = SendSummary()
        channel_id = (settings.get("active_channel") or "email").strip().lower()
        channel_obj = get_channel(channel_id)
        contacts = self.db.fetch_all(
            """
            SELECT *
            FROM contacts
            WHERE campaign_id = ? AND status = 'approved' AND channel = ?
            ORDER BY id DESC
            """,
            (campaign_id, channel_obj.channel_id),
        )

        if send_mode == "live" and channel_obj.channel_id == "telegram" and not self.telegram_bot_token():
            message = "Telegram live send blocked: сохраните Bot Token в Аккаунты и настройки → Telegram."
            summary.blocked_by_guardrail = True
            summary.errors.append(message)
            self.db.log_send(None, campaign_id, "live_send_blocked", "blocked", message, channel="telegram")
            return summary

        if send_mode == "live" and not channel_obj.supports_live_send:
            message = channel_obj.live_disabled_message
            summary.blocked_by_guardrail = True
            summary.errors.append(message)
            self.db.log_send(None, campaign_id, "live_send_blocked", "blocked", message, channel=channel_obj.channel_id)
            return summary

        if send_mode == "live" and real_send_confirm_required and not confirm_live_send:
            message = "Live send is blocked until the operator explicitly confirms."
            summary.blocked_by_guardrail = True
            summary.errors.append(message)
            self.db.log_send(None, campaign_id, "live_send_blocked", "blocked", message, channel=channel_obj.channel_id)
            return summary

        if send_mode == "live" and channel_obj.channel_id == "email" and not sender_email:
            message = "Live Gmail send is blocked: выберите активный Gmail-профиль отправителя."
            summary.blocked_by_guardrail = True
            summary.errors.append(message)
            self.db.log_send(None, campaign_id, "live_send_blocked", "blocked", message, channel="email")
            return summary

        active_profile = self.active_gmail_profile()
        if send_mode == "live" and channel_obj.channel_id == "email" and active_profile and not sender_password:
            message = "Live Gmail send is blocked: у активного Gmail-профиля нет сохраненного App Password."
            summary.blocked_by_guardrail = True
            summary.errors.append(message)
            self.db.log_send(None, campaign_id, "live_send_blocked", "blocked", message, channel="email")
            return summary

        if send_mode == "live":
            for contact in contacts:
                try:
                    self.validate_live_recipient(
                        contact,
                        safe_mode=safe_mode,
                        allowed_test_recipient=allowed_test_recipient,
                    )
                except ValueError as exc:
                    message = str(exc)
                    summary.blocked_by_guardrail = True
                    summary.errors.append(message)
                    self.db.log_send(
                        contact["id"],
                        campaign_id,
                        "send_email",
                        "blocked",
                        message,
                        channel=channel_obj.channel_id,
                        platform_recipient=self.platform_recipient(contact),
                    )
                    return summary

        if send_mode == "live" and safe_mode and daily_limit <= 0:
            message = "Safe mode requires daily_send_limit greater than 0."
            summary.errors.append(message)
            self.db.log_send(None, campaign_id, "send_email", "blocked", message)
            return summary

        for index, contact in enumerate(contacts):
            contact_id = int(contact["id"])
            if self.blacklist.is_blocked(contact["email"]):
                self.db.update_contact(
                    contact_id,
                    {"status": "blacklisted", "last_error": "Email is in blacklist"},
                )
                self.db.log_send(
                    contact_id,
                    campaign_id,
                    "send_email",
                    "blacklisted",
                    "Email is in blacklist",
                )
                summary.blacklisted += 1
                continue

            if send_mode == "dry_run":
                platform_recipient = self.validate_channel_recipient(contact)
                self.db.update_contact(
                    contact_id,
                    {
                        "status": "dry_run_sent",
                        "last_error": "",
                    },
                )
                self.db.log_send(
                    contact_id,
                    campaign_id,
                    "dry_run_send",
                    "dry_run_sent",
                    (
                        f"Dry-run only. No live message was sent. recipient_email={platform_recipient}"
                        if channel_obj.channel_id == "email"
                        else f"Dry-run only. No live message was sent. platform_recipient={platform_recipient}"
                    ),
                    channel=channel_obj.channel_id,
                    platform_recipient=platform_recipient,
                )
                self.timeline.record_event(
                    contact_id,
                    "dry_run",
                    title="Dry-run выполнен",
                    details=f"channel={channel_obj.channel_id}; recipient={platform_recipient}",
                )
                self.conversations.add_message(
                    contact_id,
                    direction="outbound",
                    subject=str(contact.get("subject") or ""),
                    body=str(contact.get("generated_message") or contact.get("base_message") or ""),
                    message_type="dry_run",
                    status="dry_run_sent",
                    metadata={"recipient": platform_recipient, "channel": channel_obj.channel_id},
                )
                summary.dry_run_sent += 1
                continue

            rate_state = self.rate_limiter.state(daily_limit)
            if not rate_state.allowed:
                summary.blocked_by_limit = True
                message = (
                    "Daily send limit reached. "
                    f"Remaining today: {rate_state.remaining}/{rate_state.daily_limit} "
                    f"(sent today: {rate_state.sent_today}/{rate_state.daily_limit})."
                )
                summary.errors.append(message)
                self.db.log_send(contact_id, campaign_id, "send_email", "blocked", message)
                break

            try:
                recipient = self.validate_live_recipient(
                    contact,
                    safe_mode=safe_mode,
                    allowed_test_recipient=allowed_test_recipient,
                )
                if channel_obj.channel_id == "telegram":
                    TelegramChannel().send_message(
                        bot_token=self.telegram_bot_token(),
                        chat_id=recipient,
                        text=contact.get("generated_message") or contact.get("base_message") or "",
                    )
                else:
                    self.mailer.send_email(
                        host=host,
                        port=port,
                        sender_email=sender_email,
                        recipient_email=recipient,
                        subject=contact.get("subject") or "",
                        body=contact.get("generated_message") or contact.get("base_message") or "",
                        password=sender_password or None,
                    )
            except ValueError as exc:
                error = redact_secret(exc)
                self.db.update_contact(contact_id, {"last_error": error})
                self.db.log_send(contact_id, campaign_id, "send_email", "blocked", error)
                summary.errors.append(error)
                summary.blocked_by_guardrail = True
                break
            except MailerConfigError as exc:
                error = redact_secret(exc)
                self.db.update_contact(contact_id, {"last_error": error})
                self.db.log_send(contact_id, campaign_id, "send_email", "blocked", error)
                summary.errors.append(error)
                break
            except MailerError as exc:
                error = redact_secret(exc)
                self.db.update_contact(contact_id, {"status": "failed", "last_error": error})
                self.db.log_send(contact_id, campaign_id, "send_email", "failed", error)
                self.timeline.record_event(contact_id, "failed", title="Ошибка", details=error)
                self.logger.error("Send failed for %s: %s", contact["email"], error)
                summary.failed += 1
                summary.errors.append(f"{contact['email']}: {error}")
                continue
            except ChannelSafetyError as exc:
                error = redact_secret(exc)
                self.db.update_contact(contact_id, {"status": "failed", "last_error": error})
                self.db.log_send(
                    contact_id,
                    campaign_id,
                    "send_message",
                    "failed",
                    error,
                    channel=channel_obj.channel_id,
                    platform_recipient=contact.get("external_id") or contact.get("handle") or "",
                )
                self.timeline.record_event(contact_id, "failed", title="Ошибка", details=error)
                self.logger.error("Channel send failed for contact %s: %s", contact_id, error)
                summary.failed += 1
                summary.errors.append(f"{contact.get('email') or contact_id}: {error}")
                continue

            sent_at = datetime.now().replace(microsecond=0)
            follow_up_due_at = sent_at + timedelta(days=follow_up_days)
            self.db.update_contact(
                contact_id,
                {
                    "status": "sent",
                    "last_error": "",
                    "sent_at": sent_at.isoformat(sep=" "),
                    "follow_up_due_at": follow_up_due_at.isoformat(sep=" "),
                },
            )
            self.db.log_send(
                contact_id,
                campaign_id,
                "send_email",
                "sent",
                (
                    f"telegram_chat_id={recipient}"
                    if channel_obj.channel_id == "telegram"
                    else f"recipient_email={recipient}"
                ),
                channel=channel_obj.channel_id,
                platform_recipient=recipient,
            )
            self.timeline.record_event(
                contact_id,
                "sent",
                title="Отправлено",
                details=f"channel={channel_obj.channel_id}; recipient={recipient}",
            )
            self.conversations.add_message(
                contact_id,
                direction="outbound",
                subject=str(contact.get("subject") or ""),
                body=str(contact.get("generated_message") or contact.get("base_message") or ""),
                message_type="sent",
                status="sent",
                metadata={"recipient": recipient, "channel": channel_obj.channel_id},
            )
            self.followups.schedule_followup(
                contact_id,
                due_at=follow_up_due_at.isoformat(sep=" "),
                note="Suggested after successful send. Manual confirmation required.",
            )
            summary.sent += 1

            if delay_seconds and index < len(contacts) - 1:
                self.sleep_fn(delay_seconds)

        return summary

    def due_followups(self, campaign_id: int) -> list[dict[str, Any]]:
        return self.db.fetch_all(
            """
            SELECT *
            FROM contacts
            WHERE campaign_id = ?
              AND status = 'sent'
              AND follow_up_due_at IS NOT NULL
              AND datetime(follow_up_due_at) <= datetime('now', 'localtime')
            ORDER BY follow_up_due_at ASC
            """,
            (campaign_id,),
        )

    def export_report(self, campaign_id: int) -> Path:
        self.export_dir.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        path = self.export_dir / f"contacts_export_{timestamp}.xlsx"

        contacts = self.db.list_contacts(campaign_id)
        send_logs = self.db.fetch_all(
            "SELECT * FROM send_logs WHERE campaign_id = ? ORDER BY created_at DESC, id DESC",
            (campaign_id,),
        )
        queue_jobs = self.db.fetch_all(
            "SELECT * FROM job_queue WHERE campaign_id = ? ORDER BY updated_at DESC, id DESC",
            (campaign_id,),
        )
        conversation_threads = self.db.fetch_all(
            """
            SELECT *
            FROM conversation_threads
            WHERE campaign_id = ?
            ORDER BY last_activity_at DESC, id DESC
            """,
            (campaign_id,),
        )
        conversation_messages = self.db.fetch_all(
            """
            SELECT *
            FROM conversation_messages
            WHERE campaign_id = ?
            ORDER BY created_at DESC, id DESC
            """,
            (campaign_id,),
        )
        errors = self.db.fetch_all(
            """
            SELECT * FROM import_errors
            WHERE campaign_id = ?
            ORDER BY created_at DESC, id DESC
            """,
            (campaign_id,),
        )

        workbook = Workbook()
        contacts_sheet = workbook.active
        contacts_sheet.title = "contacts"
        self._append_rows(contacts_sheet, contacts)

        logs_sheet = workbook.create_sheet("send_logs")
        self._append_rows(logs_sheet, send_logs)

        queue_sheet = workbook.create_sheet("job_queue")
        self._append_rows(queue_sheet, queue_jobs)

        threads_sheet = workbook.create_sheet("conversation_threads")
        self._append_rows(threads_sheet, conversation_threads)

        messages_sheet = workbook.create_sheet("conversation_messages")
        self._append_rows(messages_sheet, conversation_messages)

        errors_sheet = workbook.create_sheet("errors")
        self._append_rows(errors_sheet, errors)

        workbook.save(path)
        self.db.log_send(None, campaign_id, "export_report", "ok", str(path.name))
        return path

    @staticmethod
    def _append_rows(sheet, rows: list[dict[str, Any]]) -> None:
        if not rows:
            sheet.append(["empty"])
            return
        headers = list(rows[0].keys())
        sheet.append(headers)
        for row in rows:
            sheet.append([row.get(header, "") for header in headers])
